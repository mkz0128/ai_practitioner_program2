"""Generate deterministic F1 workbook fixtures for column mapping and validation."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).parents[1]
SOURCE = ROOT / "data" / "samples" / "demo-50-relaxed.xlsx"
OUTPUT = ROOT / "data" / "samples"


def rename_headers(workbook) -> None:
    aliases = {
        "orders": {"order_id": "訂單編號", "location_label": "收件區", "time_slot": "時段"},
        "packages": {"weight_kg": "重量kg"},
    }
    for sheet_name, field_aliases in aliases.items():
        sheet = workbook[sheet_name]
        for _index, cell in enumerate(sheet[1], start=1):
            cell.value = field_aliases.get(cell.value, cell.value)


def save_workbook(workbook, path: Path) -> None:
    workbook.save(path)


def main() -> None:
    mapped = load_workbook(SOURCE)
    rename_headers(mapped)
    save_workbook(mapped, OUTPUT / "demo-mapped-50.xlsx")

    missing = load_workbook(SOURCE)
    missing["orders"]["E2"] = None
    missing["orders"]["H3"] = None
    missing["packages"]["C4"] = None
    save_workbook(missing, OUTPUT / "demo-missing-fields.xlsx")

    duplicate = load_workbook(SOURCE)
    duplicate["orders"]["A3"] = duplicate["orders"]["A2"].value
    save_workbook(duplicate, OUTPUT / "demo-duplicate-id.xlsx")

    empty = load_workbook(SOURCE)
    for sheet in empty.worksheets:
        if sheet.max_row >= 2:
            sheet.delete_rows(2, sheet.max_row - 1)
    save_workbook(empty, OUTPUT / "demo-empty.xlsx")


if __name__ == "__main__":
    main()
