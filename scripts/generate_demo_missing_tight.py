"""Build the demo's first-act workbook: demo-50-tight with three cells blanked.

The dispatcher is shown exactly which cells are blank and types the values back
in the chat. Typing the original values reproduces demo-50-tight exactly, so the
rest of the demo's numbers are the ones that workbook has always produced.

    python scripts/generate_demo_missing_tight.py
"""

from pathlib import Path

from openpyxl import load_workbook

SOURCE = Path("data/samples/demo-50-tight.xlsx")
TARGET = Path("data/samples/demo-50-tight-missing.xlsx")

BLANKS = (
    ("orders", "order_id", "ORD-001", "location_label"),
    ("orders", "order_id", "ORD-002", "time_slot"),
    ("packages", "package_id", "PKG-003-01", "weight_kg"),
)


def main() -> None:
    workbook = load_workbook(SOURCE)
    print(f"read {SOURCE}")
    for sheet_name, id_field, record_id, field in BLANKS:
        sheet = workbook[sheet_name]
        headers = [cell.value for cell in sheet[1]]
        id_column = headers.index(id_field) + 1
        value_column = headers.index(field) + 1
        for row in range(2, sheet.max_row + 1):
            if str(sheet.cell(row=row, column=id_column).value).strip() != record_id:
                continue
            original = sheet.cell(row=row, column=value_column).value
            sheet.cell(row=row, column=value_column).value = None
            print(f"  blanked {sheet_name}.{record_id}.{field} (was {original!r})")
            break
        else:
            raise SystemExit(f"{record_id} not found in {sheet_name}")
    workbook.save(TARGET)
    print(f"wrote {TARGET}")


if __name__ == "__main__":
    main()
