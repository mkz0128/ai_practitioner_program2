from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

from agents.exceptions import ModelTimeoutError
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.agent.mapping import MappingSuggestion, MappingSuggestionOutput
from src.api import main as main_module
from src.api.main import app, saved_mapping_profiles
from src.services.importer import parse_workbook

SAMPLE_WORKBOOK = Path(__file__).parents[1] / "data" / "samples" / "demo-50-relaxed.xlsx"
MAPPED_FILENAME = "供應商-A-配送單.xlsx"
client = TestClient(app)


def mapped_workbook_bytes() -> bytes:
    workbook = load_workbook(SAMPLE_WORKBOOK)
    aliases = {
        "orders": {"order_id": "訂單編號", "location_label": "收件區", "time_slot": "時段"},
        "packages": {"weight_kg": "重量kg"},
    }
    for sheet_name, field_aliases in aliases.items():
        sheet = workbook[sheet_name]
        headers = [cell.value for cell in sheet[1]]
        for index, header in enumerate(headers, start=1):
            sheet.cell(1, index).value = field_aliases.get(header, header)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def mapping_for_mapped_workbook() -> dict[str, dict[str, str]]:
    return {
        "orders": {"訂單編號": "order_id", "收件區": "location_label", "時段": "time_slot"},
        "packages": {"重量kg": "weight_kg"},
    }


def test_importer_applies_confirmed_mapping_before_strict_models() -> None:
    dataset, report = parse_workbook(
        BytesIO(mapped_workbook_bytes()),
        source_filename=MAPPED_FILENAME,
        column_mapping=mapping_for_mapped_workbook(),
    )

    assert report.is_valid, report.model_dump()
    assert dataset is not None
    assert len(dataset.orders) == 50
    assert len(dataset.packages) == 99
    assert round(sum(order.total_weight_kg for order in dataset.orders), 3) == 316.0


def test_inspect_returns_strict_agent_mapping_suggestions(monkeypatch) -> None:
    saved_mapping_profiles.clear()

    async def fake_propose_mapping(headers_by_sheet, canonical_by_sheet):
        assert "訂單編號" in headers_by_sheet["orders"]
        assert "weight_kg" in canonical_by_sheet["packages"]
        return MappingSuggestionOutput(
            suggestions=[
                MappingSuggestion(
                    sheet="orders", source="訂單編號", target="order_id", confidence=0.98
                ),
                MappingSuggestion(
                    sheet="orders", source="收件區", target="location_label", confidence=0.95
                ),
                MappingSuggestion(
                    sheet="orders", source="時段", target="time_slot", confidence=0.93
                ),
                MappingSuggestion(
                    sheet="packages", source="重量kg", target="weight_kg", confidence=0.99
                ),
            ]
        )

    monkeypatch.setattr(main_module, "propose_mapping", fake_propose_mapping)
    response = client.post(
        "/api/v1/datasets/inspect-excel",
        files={
            "file": (
                MAPPED_FILENAME,
                mapped_workbook_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["status"] == "NEEDS_CONFIRMATION"
    assert payload["requires_confirmation"] is True
    assert payload["missing_fields"] == []
    assert payload["mapping"]["orders"]["訂單編號"] == "order_id"
    mapped_entry = next(item for item in payload["entries"] if item["source"] == "收件區")
    assert mapped_entry["confidence"] == 0.95


def test_inspect_keeps_provider_fallback_in_manual_confirmation_state(monkeypatch) -> None:
    saved_mapping_profiles.clear()

    async def fail_mapping(headers_by_sheet, canonical_by_sheet):
        raise ModelTimeoutError(1.0)

    monkeypatch.setattr(main_module, "propose_mapping", fail_mapping)
    response = client.post(
        "/api/v1/datasets/inspect-excel",
        files={
            "file": (
                MAPPED_FILENAME,
                mapped_workbook_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers={"X-Dispatch-UI": "true"},
    )

    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["status"] == "NEEDS_CONFIRMATION"
    assert payload["requires_confirmation"] is True
    assert payload["provider_fallback"]["provider"] == "OPENAI"
    assert {item["confidence"] for item in payload["entries"]} == {0.0}


def test_confirmed_mapping_import_can_be_reused_by_source_name() -> None:
    saved_mapping_profiles.clear()
    payload = mapping_for_mapped_workbook()
    response = client.post(
        "/api/v1/datasets/import-excel",
        files={
            "file": (
                MAPPED_FILENAME,
                mapped_workbook_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"mapping": json.dumps(payload, ensure_ascii=False), "mapping_name": "供應商 A"},
    )

    assert response.status_code == 201, response.text
    assert MAPPED_FILENAME in saved_mapping_profiles
    assert "供應商 A" in saved_mapping_profiles

    inspected = client.post(
        "/api/v1/datasets/inspect-excel",
        files={
            "file": (
                MAPPED_FILENAME,
                mapped_workbook_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert inspected.status_code == 200, inspected.text
    assert inspected.json()["status"] == "AUTO_APPLIED"
    assert inspected.json()["requires_confirmation"] is False


def test_importer_reports_all_required_cell_errors_in_one_report() -> None:
    workbook = load_workbook(SAMPLE_WORKBOOK)
    workbook["orders"]["E2"] = None
    workbook["orders"]["H3"] = None
    workbook["packages"]["C4"] = None
    output = BytesIO()
    workbook.save(output)

    _dataset, report = parse_workbook(BytesIO(output.getvalue()))

    assert not report.is_valid
    paths = {error.path for error in report.errors}
    assert "orders.ORD-001.location_label" in paths
    assert "orders.ORD-002.time_slot" in paths
    assert "packages.PKG-003-01.weight_kg" in paths


def test_importer_rejects_header_only_workbook() -> None:
    workbook = load_workbook(SAMPLE_WORKBOOK)
    for sheet in workbook.worksheets:
        sheet.delete_rows(2, sheet.max_row)
    output = BytesIO()
    workbook.save(output)

    dataset, report = parse_workbook(BytesIO(output.getvalue()))

    assert dataset is None
    assert report.errors[0].code == "EMPTY_DATASET"
    assert "至少提供一筆訂單" in report.errors[0].message


def test_ui_inspection_keeps_expected_validation_errors_out_of_browser_console() -> None:
    payload = {
        "file": (
            "bad.xlsx",
            b"not an excel workbook",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }

    direct = client.post("/api/v1/datasets/inspect-excel", files=payload)
    ui = client.post(
        "/api/v1/datasets/inspect-excel",
        files=payload,
        headers={"X-Dispatch-UI": "true"},
    )

    assert direct.status_code == 400
    assert direct.json()["error"]["code"] == "INVALID_XLSX"
    assert ui.status_code == 200
    assert ui.json()["status"] == "INVALID"
    assert ui.json()["error"]["message"] == "無法讀取 .xlsx 檔案。"
