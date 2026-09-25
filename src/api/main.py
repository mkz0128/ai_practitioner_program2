from __future__ import annotations

import hashlib
import hmac
import inspect
import json
import re
from dataclasses import dataclass, field, replace
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from threading import RLock
from typing import Annotated, Any, Literal, cast
from uuid import uuid4

from agents import InputGuardrailTripwireTriggered
from agents.exceptions import ModelBehaviorError, ModelTimeoutError, ToolTimeoutError
from fastapi import FastAPI, File, Form, Request, UploadFile
from fastapi.exceptions import RequestValidationError
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from openai import (
    APIConnectionError,
    APITimeoutError,
    AuthenticationError,
    BadRequestError,
    PermissionDeniedError,
    RateLimitError,
)
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pydantic import BaseModel, ConfigDict, Field

from src.agent.mapping import positional_mapping, propose_mapping
from src.agent.runtime import run_dispatch_agent
from src.agent.urgent_workflow import (
    UrgentOrderDraft,
    UrgentUnderstanding,
    UrgentWorkflowState,
    advance_urgent_workflow,
    understand_urgent_message,
)
from src.config import get_settings
from src.domain.models import (
    Dataset,
    Order,
    Package,
    Priority,
    TimeSlotValue,
    Vehicle,
    VehicleStatus,
    Zone,
)
from src.providers.google_routes import GoogleRoutesProvider, GoogleRoutesProviderError
from src.providers.tdx import TDXProvider, correlate_events_to_plan
from src.repositories.sqlite import SQLiteRepository
from src.services.demo_orders import get_demo_urgent_order
from src.services.dispatch_deviations import compute_dispatch_deviations
from src.services.dispatch_parameters import (
    _parameterized_route,
    apply_service_time_parameters,
    confirm_service_parameter,
    parameter_state,
    reset_service_parameters,
)
from src.services.dispatch_rules import (
    DispatchRule,
    DispatchRuleDraft,
    build_plan_with_rules,
    deactivate_dispatch_rule,
    expires_at_for_duration,
    list_dispatch_rules,
    reset_dispatch_rules,
    rule_summary,
    save_dispatch_rule,
)
from src.services.dispatch_rules import (
    preview_dispatch_rule as preview_rule_trial,
)
from src.services.display import slot_sentence, vehicle_label
from src.services.errors import ValidationReport
from src.services.evidence import recommendation_reason
from src.services.fingerprint import dataset_hash, matrix_hash
from src.services.importer import SHEET_FIELDS, ColumnMapping, parse_workbook, validate_dataset
from src.services.matrix import MatrixResult, SimulatedRouteProvider
from src.services.plan_diff import compute_plan_diff
from src.services.planner import (
    PREFERRED_VEHICLE_BY_ZONE,
    Objective,
    PlanResult,
    build_baseline,
    build_ortools,
    preview_reassignment,
    uses_legacy_timing,
)
from src.services.risk import calculate_plan_risks, summarize_delay
from src.services.solve_scope import (
    clamp_timeline_minutes,
    route_progress,
    scope_payload,
    stage_for_state,
)
from src.services.urgent_options import (
    UrgentOption,
    build_partial_urgent_plan,
    build_urgent_options,
    count_reordered_orders,
)
from src.services.validator import PlanValidation, validate_plan


class StrictRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")


def _classify_agent_error(exc: Exception) -> tuple[int, str, str, bool]:
    """Return a safe client-facing classification without serializing SDK details."""
    if isinstance(exc, RuntimeError) and str(exc) == "AGENT_DID_NOT_CALL_PLAN_TOOL":
        return 502, "AGENT_TOOL_SELECTION_FAILED", "AI 助理未能完成必要的工具選擇，請重試。", True
    if isinstance(exc, (ModelTimeoutError, ToolTimeoutError, APITimeoutError)):
        return 504, "AGENT_TIMEOUT", "AI 助理回應逾時，請稍後重試。", True
    if isinstance(exc, RateLimitError):
        return 503, "AGENT_RATE_LIMITED", "AI 服務目前忙碌，請稍後重試。", True
    if isinstance(exc, APIConnectionError):
        return 503, "AGENT_PROVIDER_UNAVAILABLE", "AI 服務目前無法連線，請稍後重試。", True
    if isinstance(exc, (AuthenticationError, PermissionDeniedError)):
        return 503, "AGENT_CREDENTIALS_REJECTED", "AI 服務授權失敗，請由管理者檢查設定。", False
    if isinstance(exc, BadRequestError):
        return 502, "AGENT_REQUEST_REJECTED", "AI 服務無法處理這次要求，請調整內容後重試。", False
    if isinstance(exc, ModelBehaviorError):
        return 502, "AGENT_INVALID_RESPONSE", "AI 回覆未通過安全檢查，方案沒有變更。", False
    return 502, "AGENT_RUN_FAILED", "AI 助理暫時無法完成這次要求，方案沒有變更。", False


class CreatePlanRequest(StrictRequest):
    dataset_id: str
    route_provider_preference: Literal["AUTO", "SIMULATED"] = "AUTO"
    traffic_mode: Literal["AUTO", "SIMULATED"] = "AUTO"
    simulation_seed: int = 20260901
    algorithm: Literal["BASELINE", "ORTOOLS"] = "ORTOOLS"
    objective: Objective = "FASTEST"


class ConfirmRequest(StrictRequest):
    version: int = Field(ge=1)
    confirmation: Literal["CONFIRM_PLAN"]
    dispatcher_reference: str = Field(min_length=1, max_length=120)


class DispatchRequest(StrictRequest):
    version: int = Field(ge=1)
    confirmation: Literal["MARK_DISPATCHED"]


class LoadingRequest(StrictRequest):
    version: int = Field(ge=1)
    confirmation: Literal["START_LOADING"]
    dispatcher_reference: str = Field(min_length=1, max_length=120)


class SimulatedDepartureRequest(StrictRequest):
    version: int = Field(ge=1)
    confirmation: Literal["START_SIMULATED_DEPARTURE"]


class CompareStrategiesRequest(StrictRequest):
    dataset_id: str
    plan_id: str | None = None
    version: int | None = Field(default=None, ge=1)
    route_provider_preference: Literal["AUTO", "SIMULATED"] = "AUTO"
    traffic_mode: Literal["AUTO", "SIMULATED"] = "AUTO"


class DelaySimulationRequest(StrictRequest):
    version: int = Field(ge=1)
    delay_minutes: Literal[10, 20, 30]


class ReassignmentRequest(StrictRequest):
    base_plan_version: int = Field(ge=1)
    order_id: str = Field(min_length=1)
    target_vehicle_id: str = Field(min_length=1)


class RouteOrderRequest(StrictRequest):
    base_plan_version: int = Field(ge=1)
    vehicle_id: str = Field(min_length=1)
    order_ids: list[str] = Field(max_length=200)
    timeline_minutes: int | None = Field(default=None, ge=0, le=660)


class CrossVehicleRouteOrderRequest(StrictRequest):
    base_plan_version: int = Field(ge=1)
    source_vehicle_id: str = Field(min_length=1)
    target_vehicle_id: str = Field(min_length=1)
    order_id: str = Field(min_length=1)
    target_sequence: int = Field(ge=1)
    timeline_minutes: int | None = Field(default=None, ge=0, le=660)


class RestorePlanRequest(StrictRequest):
    source_version: int = Field(ge=1)
    dispatcher_reference: str = Field(min_length=1, max_length=120)


class ChatRequest(StrictRequest):
    session_id: str = Field(min_length=1, max_length=120)
    message: str = Field(min_length=1, max_length=4000)
    context: dict[str, Any] = Field(default_factory=dict)
    action: Literal["PREVIEW_URGENT"] | None = None


class DemoLoginRequest(StrictRequest):
    password: str = Field(min_length=1, max_length=256)


class UrgentOrderRequest(StrictRequest):
    order_id: str
    zone_code: str
    city: str
    district: str
    location_label: str
    latitude: float
    longitude: float
    time_slot: TimeSlotValue
    declared_package_count: int = Field(ge=1, le=3)
    priority: Priority = Priority.NORMAL
    note: str | None = None


class UrgentInsertRequest(StrictRequest):
    base_plan_version: int = Field(ge=1)
    order: UrgentOrderRequest
    packages: list[Package] = Field(min_length=1, max_length=3)


class UrgentOrderBundleRequest(StrictRequest):
    order: UrgentOrderRequest
    packages: list[Package] = Field(min_length=1, max_length=20)


class UrgentBatchInsertRequest(StrictRequest):
    base_plan_version: int = Field(ge=1)
    orders: list[UrgentOrderBundleRequest] = Field(min_length=1, max_length=20)


class ColumnMappingRequest(StrictRequest):
    mapping: dict[str, dict[str, str]]
    source_name: str | None = None
    save_as: str | None = Field(default=None, max_length=120)


class AdditionalDispatchRuleRequest(StrictRequest):
    rule_type: Literal["LATEST_RETURN_TIME"]
    value: str
    duration: Literal["PERMANENT", "THIS_WEEK", "TODAY"] = "PERMANENT"


class DispatchRuleMutationRequest(StrictRequest):
    plan_id: str = Field(min_length=1)
    base_plan_version: int = Field(ge=1)
    subject_type: Literal["VEHICLE", "ZONE"] = "VEHICLE"
    subject_id: str = Field(min_length=1)
    rule_type: Literal[
        "MAX_PACKAGE_WEIGHT",
        "MAX_ROUTE_DISTANCE",
        "MAX_STOPS",
        "EXCLUDED_ZONE",
        "ALLOWED_TIME_WINDOW",
        "LATEST_RETURN_TIME",
    ]
    value: float | str
    source_utterance: str = Field(min_length=1, max_length=4000)
    duration: Literal["PERMANENT", "THIS_WEEK", "TODAY"] = "PERMANENT"
    additional_rule: AdditionalDispatchRuleRequest | None = None


class DispatchParameterConfirmationRequest(StrictRequest):
    plan_id: str = Field(min_length=1)
    base_plan_version: int = Field(ge=1)
    zone_code: str = Field(min_length=1)
    from_service_minutes: int = Field(ge=1, le=15)
    to_service_minutes: int = Field(ge=1, le=15)
    source: Literal["TIMELINE_DEVIATION"] = "TIMELINE_DEVIATION"


@dataclass
class DatasetRecord:
    dataset_id: str
    dataset: Dataset
    validation: ValidationReport
    matrix: MatrixResult
    created_at: str


@dataclass
class PlanRecord:
    plan_id: str
    dataset_id: str
    version: int
    state: str
    plan: PlanResult
    validation: PlanValidation
    matrix: MatrixResult
    created_at: str


@dataclass
class AgentSession:
    """Structured conversation pointers; secrets and workbook payloads are never retained."""

    dataset_id: str | None = None
    plan_id: str | None = None
    plan_version: int | None = None
    order_id: str | None = None
    vehicle_id: str | None = None
    strategy: str | None = None
    rule_source_utterance: str | None = None
    frozen_stop_count: int = 0
    frozen_stop_ids: tuple[str, ...] = ()
    pending_fields: tuple[str, ...] = ()
    last_preview_version: int | None = None
    last_tool: str | None = None
    pending_order: dict[str, Any] | None = None
    urgent_workflow: dict[str, Any] = field(default_factory=dict)
    history: list[tuple[str, str]] = field(default_factory=list)


_SESSION_SECRET_PATTERN = re.compile(
    r"(?i)(api[_ -]?key|client[_ -]?secret|access[_ -]?token|password)\s*[=:]\s*[^\s,;]+"
)


def _safe_session_text(value: str) -> str:
    """Keep conversational context useful without retaining credential values."""
    return _SESSION_SECRET_PATTERN.sub(r"\1=[REDACTED]", value)[:4000]


def _session_payload(session: AgentSession) -> dict[str, Any]:
    return {
        "dataset_id": session.dataset_id,
        "plan_id": session.plan_id,
        "plan_version": session.plan_version,
        "order_id": session.order_id,
        "vehicle_id": session.vehicle_id,
        "strategy": session.strategy,
        "rule_source_utterance": session.rule_source_utterance,
        "frozen_stop_count": session.frozen_stop_count,
        "frozen_stop_ids": list(session.frozen_stop_ids),
        "pending_fields": list(session.pending_fields),
        "last_preview_version": session.last_preview_version,
        "last_tool": session.last_tool,
        "pending_order": session.pending_order,
        "urgent_workflow": session.urgent_workflow,
        "history": [[role, _safe_session_text(content)] for role, content in session.history[-12:]],
    }


def _session_from_payload(payload: dict[str, Any]) -> AgentSession:
    raw_history = payload.get("history", [])
    history = [
        (str(item[0]), _safe_session_text(str(item[1])))
        for item in raw_history
        if isinstance(item, list) and len(item) == 2
    ]
    pending = payload.get("pending_fields", [])
    plan_id = payload.get("plan_id")
    plan_version = payload.get("plan_version")
    order_id = payload.get("order_id")
    vehicle_id = payload.get("vehicle_id")
    strategy = payload.get("strategy")
    rule_source_utterance = payload.get("rule_source_utterance")
    frozen_stop_count = payload.get("frozen_stop_count")
    frozen_stop_ids = payload.get("frozen_stop_ids")
    last_preview_version = payload.get("last_preview_version")
    last_tool = payload.get("last_tool")
    pending_order = payload.get("pending_order")
    urgent_workflow = payload.get("urgent_workflow")
    return AgentSession(
        dataset_id=(
            payload.get("dataset_id") if isinstance(payload.get("dataset_id"), str) else None
        ),
        plan_id=plan_id if isinstance(plan_id, str) else None,
        plan_version=plan_version if isinstance(plan_version, int) else None,
        order_id=order_id if isinstance(order_id, str) else None,
        vehicle_id=vehicle_id if isinstance(vehicle_id, str) else None,
        strategy=strategy if isinstance(strategy, str) else None,
        rule_source_utterance=(
            rule_source_utterance if isinstance(rule_source_utterance, str) else None
        ),
        frozen_stop_count=frozen_stop_count if isinstance(frozen_stop_count, int) else 0,
        frozen_stop_ids=(
            tuple(item for item in frozen_stop_ids if isinstance(item, str))
            if isinstance(frozen_stop_ids, list)
            else ()
        ),
        pending_fields=(
            tuple(item for item in pending if isinstance(item, str))
            if isinstance(pending, list)
            else ()
        ),
        last_preview_version=(
            last_preview_version if isinstance(last_preview_version, int) else None
        ),
        last_tool=last_tool if isinstance(last_tool, str) else None,
        pending_order=pending_order if isinstance(pending_order, dict) else None,
        urgent_workflow=urgent_workflow if isinstance(urgent_workflow, dict) else {},
        history=history[-12:],
    )


@dataclass
class InMemoryStore:
    datasets: dict[str, DatasetRecord] = field(default_factory=dict)
    plans: dict[str, dict[int, PlanRecord]] = field(default_factory=dict)
    current_versions: dict[str, int] = field(default_factory=dict)
    lock: RLock = field(default_factory=RLock)

    def add_dataset(self, record: DatasetRecord) -> None:
        with self.lock:
            self.datasets[record.dataset_id] = record

    def add_plan(self, record: PlanRecord, make_current: bool = True) -> None:
        with self.lock:
            self.plans.setdefault(record.plan_id, {})[record.version] = record
            if make_current:
                self.current_versions[record.plan_id] = record.version

    def get_dataset(self, dataset_id: str) -> DatasetRecord | None:
        with self.lock:
            return self.datasets.get(dataset_id)

    def get_plan(self, plan_id: str, version: int | None = None) -> PlanRecord | None:
        with self.lock:
            versions = self.plans.get(plan_id)
            if not versions:
                return None
            selected = (
                self.current_versions.get(plan_id, max(versions)) if version is None else version
            )
            return versions.get(selected)


def _build_matrix(dataset: Dataset, *, prefer_live: bool) -> MatrixResult:
    """Resolve the matrix once and make live failures explicit when requested."""
    if not prefer_live:
        return SimulatedRouteProvider().build(dataset)
    if not settings.google_routes_enabled:
        return replace(SimulatedRouteProvider().build(dataset), warning="GOOGLE_ROUTES_DISABLED")
    if not settings.google_routes_server_api_key:
        return replace(SimulatedRouteProvider().build(dataset), warning="GOOGLE_KEY_MISSING")
    try:
        matrix = GoogleRoutesProvider(settings.google_routes_server_api_key).build(
            dataset, allow_fallback=False
        )
    except GoogleRoutesProviderError:
        provider_runtime_state["google_routes"] = "failed"
        raise
    provider_runtime_state["google_routes"] = "connected"
    return matrix


def _dataset_matrix_coordinates(dataset: Dataset) -> list[tuple[float, float]]:
    orders = tuple(sorted(dataset.orders, key=lambda order: order.order_id))
    return [
        (SimulatedRouteProvider.depot_latitude, SimulatedRouteProvider.depot_longitude),
        *[(order.latitude, order.longitude) for order in orders],
    ]


store = InMemoryStore()
agent_sessions: dict[str, AgentSession] = {}
saved_mapping_profiles: dict[str, ColumnMapping] = {}
app = FastAPI(title="AI Delivery Dispatch Agent", version="0.1.0")
settings = get_settings()
provider_runtime_state: dict[str, str] = {
    "google_routes": "configured"
    if settings.google_routes_enabled and settings.google_routes_server_api_key
    else "disabled",
    "openai": "configured" if settings.openai_api_key else "disabled",
}
repository = SQLiteRepository(settings.database_url)
app.add_middleware(
    CORSMiddleware,
    allow_origins=settings.cors_origins,
    allow_credentials=False,
    allow_methods=["GET", "POST"],
    allow_headers=["*"],
)


@app.exception_handler(RequestValidationError)
async def request_validation_exception_handler(
    request: Request, exc: RequestValidationError
) -> JSONResponse:
    """Keep malformed API payloads in the same safe, field-level envelope.

    FastAPI's default 422 response includes raw ``input`` values and a generic
    ``detail`` array.  The API contract instead requires actionable paths and
    manual-review signalling without echoing untrusted payload contents.
    """

    field_errors: list[dict[str, Any]] = []
    for item in exc.errors():
        location = [str(part) for part in item.get("loc", ()) if part != "body"]
        path = ".".join(location) or "request"
        error_type = str(item.get("type", ""))
        code = "MISSING_REQUIRED_FIELD" if error_type == "missing" else "FIELD_VALIDATION_ERROR"
        message = (
            "缺少必要欄位，請補齊後再試。"
            if error_type == "missing"
            else "欄位格式不正確，請修正後再試。"
        )
        field_errors.append(
            {
                "path": path,
                "code": code,
                "message": message,
                "value_summary": None,
                "requires_manual_review": True,
            }
        )
    return _error(
        request,
        422,
        "FIELD_VALIDATION_ERROR",
        "資料欄位未通過驗證，請依欄位提示修正後再試。",
        field_errors=field_errors,
        requires_manual_review=True,
    )


def _demo_session_token() -> str | None:
    password = settings.demo_access_password
    if not password:
        return None
    return hmac.new(
        password.encode("utf-8"), b"ai-dispatch-demo-session", hashlib.sha256
    ).hexdigest()


def _has_demo_session(request: Request) -> bool:
    expected = _demo_session_token()
    actual = request.cookies.get("dispatch_demo_session")
    return bool(expected and actual and hmac.compare_digest(actual, expected))


@app.middleware("http")
async def demo_access_middleware(request: Request, call_next: Any) -> Any:
    """Protect mutating/data APIs when a Render demo password is configured.

    Local development and deterministic tests leave DEMO_ACCESS_PASSWORD unset,
    preserving the existing API contract. Health, docs, login, and the public
    browser-key runtime configuration remain reachable before login.
    """
    protected = request.url.path.startswith("/api/v1/") and request.url.path not in {
        "/api/v1/runtime-config",
    }
    if (
        settings.demo_access_password
        and protected
        and request.method != "OPTIONS"
        and not _has_demo_session(request)
    ):
        return JSONResponse(
            status_code=401,
            content={
                "error": {
                    "code": "DEMO_AUTH_REQUIRED",
                    "message": "請先登入展示環境。",
                    "field_errors": [],
                }
            },
        )
    return await call_next(request)


def _request_id(request: Request) -> str:
    return getattr(request.state, "request_id", f"REQ-{uuid4().hex[:12]}")


def _empty_agent_dataset() -> tuple[Dataset, MatrixResult]:
    dataset = Dataset(orders=(), packages=(), vehicles=(), zones=())
    return dataset, SimulatedRouteProvider().build(dataset)


def _error(
    request: Request, status_code: int, code: str, message: str, **details: Any
) -> JSONResponse:
    request_id = _request_id(request)
    return JSONResponse(
        status_code=status_code,
        content={
            "error": {
                "code": code,
                "message": message,
                "field_errors": details.pop("field_errors", []),
                "request_id": request_id,
                "details": details,
            },
            "request_id": request_id,
        },
        headers={"X-Request-ID": request_id},
    )


def _validation_payload(report: ValidationReport) -> dict[str, Any]:
    return {
        "is_valid": report.is_valid,
        "error_count": len(report.errors),
        "warning_count": len(report.warnings),
        "requires_manual_review": report.requires_manual_review,
        "errors": [error.model_dump() for error in report.errors],
        "warnings": [warning.model_dump() for warning in report.warnings],
    }


def _plan_payload(record: PlanRecord) -> dict[str, Any]:
    dataset_record = store.get_dataset(record.dataset_id)
    dataset = dataset_record.dataset if dataset_record else None
    orders = {order.order_id: order for order in dataset.orders} if dataset else {}
    packages = {order_id: 0 for order_id in orders}
    for package in dataset.packages if dataset else ():
        packages[package.order_id] = packages.get(package.order_id, 0) + 1
    routes: list[dict[str, Any]] = []
    route_by_vehicle = {route.vehicle_id: route for route in record.plan.routes}
    risk_by_order = (
        {item["order_id"]: item for item in calculate_plan_risks(dataset, record.plan)}
        if dataset
        else {}
    )
    for route in record.plan.routes:
        vehicle = (
            next((item for item in dataset.vehicles if item.vehicle_id == route.vehicle_id), None)
            if dataset
            else None
        )
        cumulative_load = vehicle.current_load_kg if vehicle else 0.0
        previous_node_id = "DEPOT-001"
        stops: list[dict[str, Any]] = []
        for stop in route.stops:
            order = orders.get(stop.order_id)
            if order is None:
                continue
            cumulative_load = round(cumulative_load + order.total_weight_kg, 3)
            capacity_avoidance = None
            if dataset and vehicle:
                for candidate in sorted(dataset.vehicles, key=lambda item: item.vehicle_id):
                    if candidate.vehicle_id == vehicle.vehicle_id:
                        continue
                    if candidate.status.value != "AVAILABLE":
                        continue
                    if order.zone_code not in candidate.service_zone_codes:
                        continue
                    candidate_route = route_by_vehicle.get(candidate.vehicle_id)
                    candidate_load = (
                        candidate_route.planned_load_kg
                        if candidate_route is not None
                        else candidate.current_load_kg
                    )
                    if candidate_load + order.total_weight_kg > candidate.max_load_kg:
                        capacity_avoidance = {
                            "source_vehicle_id": candidate.vehicle_id,
                            "source_vehicle_max_load_kg": candidate.max_load_kg,
                            "assigned_vehicle_id": vehicle.vehicle_id,
                            "candidate_post_assignment_load_kg": round(
                                candidate_load + order.total_weight_kg, 3
                            ),
                        }
                        break
            stop_payload = {
                **stop.model_dump(),
                "location_label": order.location_label,
                "reason": recommendation_reason(
                    route,
                    stop,
                    vehicle,
                    order,
                    previous_node_id,
                    cumulative_load,
                    record.matrix.provider_mode,
                    record.validation.valid,
                    record.plan.algorithm,
                    capacity_avoidance,
                )
                if vehicle
                else None,
                "risk": risk_by_order.get(stop.order_id),
            }
            stops.append(stop_payload)
            previous_node_id = stop.order_id
        unused_reason = None
        if not route.order_ids:
            unused_reason = (
                "其他車輛已在不違反限制下完成全部訂單，此車保留備援容量。"
                if record.plan.complete
                else "目前沒有剩餘訂單能在載重、服務區域與時段限制內合法安排至此車。"
            )
        route_zone_codes = sorted(
            {
                orders[order_id].zone_code
                for order_id in route.order_ids
                if order_id in orders
            }
        )
        routes.append(
            {
                "vehicle_id": route.vehicle_id,
                "vehicle_name": vehicle.vehicle_name if vehicle else route.vehicle_id,
                # The card describes this plan's actual geographic
                # responsibility.  The dataset vehicle zones remain the
                # eligibility boundary used by the deterministic solver.
                "service_zone_codes": (
                    route_zone_codes
                    if route_zone_codes
                    else list(vehicle.service_zone_codes) if vehicle else []
                ),
                "order_count": len(route.order_ids),
                "package_count": sum(packages.get(order_id, 0) for order_id in route.order_ids),
                "planned_load_kg": route.planned_load_kg,
                "max_load_kg": route.max_load_kg,
                "load_utilization": route.load_utilization,
                "total_distance_m": route.total_distance_m,
                "total_duration_s": route.total_duration_s,
                "route_provider_mode": record.matrix.provider_mode,
                "unused_reason": unused_reason,
                "stops": stops,
            }
        )
    assigned = sum(len(route.order_ids) for route in record.plan.routes)
    assigned_order_ids = {order_id for route in record.plan.routes for order_id in route.order_ids}
    total_packages = sum(packages.values())
    total_weight = sum(order.total_weight_kg for order in dataset.orders) if dataset else 0.0
    assigned_weight = (
        sum(
            orders[order_id].total_weight_kg
            for order_id in assigned_order_ids
            if order_id in orders
        )
        if dataset
        else 0.0
    )
    current_dataset_hash = dataset_hash(dataset) if dataset else None
    total_orders = len(orders)
    is_complete = (
        record.plan.complete and assigned == total_orders and not record.plan.unassigned_orders
    )
    rule_check_passed = record.validation.valid
    confirmation_blockers: list[str] = []
    if record.plan.algorithm != "ORTOOLS":
        confirmation_blockers.append("NOT_FORMAL_OPTIMIZED_PLAN")
    if not is_complete:
        confirmation_blockers.append("UNASSIGNED_ORDERS")
    if not rule_check_passed:
        confirmation_blockers.append("RULE_CHECK_FAILED")
    if record.state != "PROPOSED":
        confirmation_blockers.append("PLAN_STATE_NOT_PROPOSED")
    can_confirm = not confirmation_blockers
    warnings: list[dict[str, Any]] = []
    if record.matrix.provider_mode == "SIMULATED":
        warnings.append(
            {
                "code": record.matrix.warning or "SIMULATED_ROUTE_DATA",
                "message": "目前使用可重現的模擬距離與路線資料, 非 Google 即時資料。",
            }
        )
    elif record.matrix.warning:
        warnings.append({"code": record.matrix.warning, "message": "路線 provider 回傳警告。"})
    return {
        "plan_id": record.plan_id,
        "version": record.version,
        "dataset_id": record.dataset_id,
        "state": record.state,
        "stage": stage_for_state(record.state),
        "solve_scope": scope_payload(record.plan, record.state),
        "timezone": "Asia/Taipei",
        "provider_mode": record.matrix.provider_mode,
        "matrix_hash": matrix_hash(record.matrix),
        "matrix_version": record.matrix.matrix_version,
        "algorithm": record.plan.algorithm,
        "objective": record.plan.objective,
        "dataset_hash": current_dataset_hash,
        "is_fully_feasible": record.plan.complete and record.validation.valid,
        "completeness": {
            "is_complete": is_complete,
            "assigned_order_count": assigned,
            "total_order_count": total_orders,
            "unassigned_order_count": len(record.plan.unassigned_orders),
        },
        "rule_check": {
            "passed": rule_check_passed,
            "violations": record.validation.violations,
        },
        "confirmability": {
            "can_confirm": can_confirm,
            "blockers": confirmation_blockers,
        },
        "requires_human_confirmation": True,
        "summary": {
            "assigned_order_count": assigned,
            "unassigned_order_count": len(record.plan.unassigned_orders),
            "total_package_count": total_packages,
            "total_weight_kg": round(total_weight, 3),
            "assigned_weight_kg": round(assigned_weight, 3),
            "total_distance_m": record.plan.total_distance_m,
            "total_duration_s": record.plan.total_driving_time_s,
            "algorithm": record.plan.algorithm,
            "objective": record.plan.objective,
            "dataset_hash": current_dataset_hash,
            "matrix_hash": matrix_hash(record.matrix),
            "matrix_version": record.matrix.matrix_version,
            "unassigned_orders": list(record.plan.unassigned_orders),
            "vehicles": [
                {
                    "vehicle_id": route.vehicle_id,
                    "planned_load_kg": route.planned_load_kg,
                    "max_load_kg": route.max_load_kg,
                    "load_utilization": route.load_utilization,
                }
                for route in record.plan.routes
            ],
        },
        "vehicles": routes,
        "unassigned_orders": record.plan.unassigned_orders,
        "unassigned_reasons": _readable_unassigned_reasons(dataset, record.plan),
        "validation": record.validation.model_dump(),
        "warnings": warnings,
        "parameter_state": parameter_state(),
        "parameter_replan": {
            "applied": record.plan.solver_status == "PARAMETERIZED_REPLAN",
            "message": (
                "已用已確認的區域服務時間重新計算；未能容納的訂單會明列。"
                if record.plan.solver_status == "PARAMETERIZED_REPLAN"
                else None
            ),
        },
        "created_at": record.created_at,
    }


def _route_snapshot(route: Any) -> dict[str, Any]:
    return {
        "vehicle_id": route.vehicle_id,
        "order_ids": list(route.order_ids),
        "total_distance_m": route.total_distance_m,
        "total_duration_s": route.total_duration_s,
        "planned_load_kg": route.planned_load_kg,
        "load_utilization": route.load_utilization,
        "stops": [{"order_id": stop.order_id, "eta": stop.eta} for stop in route.stops],
    }


def _route_order_candidate(
    base: PlanRecord,
    dataset_record: DatasetRecord,
    payload: RouteOrderRequest,
) -> tuple[Any, PlanResult | None, str | None, PlanValidation]:
    route = next(
        (item for item in base.plan.routes if item.vehicle_id == payload.vehicle_id),
        None,
    )
    if route is None:
        return None, None, "找不到指定車輛。", base.validation
    requested_ids = list(payload.order_ids)
    if len(requested_ids) != len(set(requested_ids)):
        return route, None, "站序不能包含重複訂單。", base.validation
    if set(requested_ids) != set(route.order_ids):
        return route, None, "只能在同一台車內調整既有站點順序。", base.validation
    if base.state == "DISPATCHED":
        progress = next(
            (item for item in route_progress(base.plan, payload.timeline_minutes)
             if item["vehicle_id"] == payload.vehicle_id),
            None,
        )
        frozen = list(progress["completed_stops"]) if progress else []
        if requested_ids[:len(frozen)] != frozen:
            return route, None, "已發車的已完成站點不能拖曳。", base.validation
    vehicle = next(
        (item for item in dataset_record.dataset.vehicles if item.vehicle_id == payload.vehicle_id),
        None,
    )
    if vehicle is None:
        return route, None, "找不到指定車輛。", base.validation
    orders = {item.order_id: item for item in dataset_record.dataset.orders}
    candidate_route = _parameterized_route(
        requested_ids,
        vehicle,
        orders,
        base.matrix,
        parameter_state()["service_minutes_by_zone"],
        uses_legacy_timing(dataset_record.dataset),
    )
    if candidate_route is None:
        return route, None, "這個站序無法同時符合配送時段或回站時間限制。", base.validation
    if candidate_route.planned_load_kg > vehicle.max_load_kg + 1e-6:
        return route, None, "這個站序會超出車輛載重上限。", base.validation
    routes = [
        candidate_route if item.vehicle_id == payload.vehicle_id else item
        for item in base.plan.routes
    ]
    candidate_plan = base.plan.model_copy(
        update={
            "routes": routes,
            "total_distance_m": sum(item.total_distance_m for item in routes),
            "total_driving_time_s": sum(item.total_duration_s for item in routes),
        }
    )
    validation = validate_plan(dataset_record.dataset, candidate_plan, base.matrix)
    if not validation.valid:
        return route, candidate_plan, "這個站序未通過配送區域、載重或時段驗證。", validation
    return route, candidate_plan, None, validation


def _cross_vehicle_route_order_candidate(
    base: PlanRecord,
    dataset_record: DatasetRecord,
    payload: CrossVehicleRouteOrderRequest,
) -> tuple[PlanResult | None, str | None, PlanValidation, Any, Any]:
    source = next(
        (item for item in base.plan.routes if item.vehicle_id == payload.source_vehicle_id),
        None,
    )
    target = next(
        (item for item in base.plan.routes if item.vehicle_id == payload.target_vehicle_id),
        None,
    )
    if source is None or target is None:
        return None, "找不到指定車輛。", base.validation, source, target
    if source.vehicle_id == target.vehicle_id:
        return None, "來源與目標車輛必須不同。", base.validation, source, target
    if payload.order_id not in source.order_ids:
        return None, "這張訂單不在來源車輛上。", base.validation, source, target
    if base.state in {"LOADED", "DISPATCHED"}:
        return None, "開始裝車後不能跨車移動；既有車輛指派已鎖定。", base.validation, source, target
    source_ids = [order_id for order_id in source.order_ids if order_id != payload.order_id]
    target_ids = list(target.order_ids)
    target_index = min(payload.target_sequence - 1, len(target_ids))
    target_ids.insert(target_index, payload.order_id)
    vehicles = {item.vehicle_id: item for item in dataset_record.dataset.vehicles}
    orders = {item.order_id: item for item in dataset_record.dataset.orders}
    service_by_zone = parameter_state()["service_minutes_by_zone"]
    legacy = uses_legacy_timing(dataset_record.dataset)
    source_vehicle = vehicles.get(source.vehicle_id)
    target_vehicle = vehicles.get(target.vehicle_id)
    if source_vehicle is None or target_vehicle is None:
        return None, "找不到指定車輛。", base.validation, source, target
    source_candidate = _parameterized_route(
        source_ids, source_vehicle, orders, base.matrix, service_by_zone, legacy
    )
    target_candidate = _parameterized_route(
        target_ids, target_vehicle, orders, base.matrix, service_by_zone, legacy
    )
    if source_candidate is None or target_candidate is None:
        return None, "這次換車會讓配送時段或回站時間不合法。", base.validation, source, target
    if target_candidate.planned_load_kg > target_vehicle.max_load_kg + 1e-6:
        return None, "這次換車會超出目標車輛載重上限。", base.validation, source, target
    routes = [
        source_candidate if item.vehicle_id == source.vehicle_id
        else target_candidate if item.vehicle_id == target.vehicle_id
        else item
        for item in base.plan.routes
    ]
    candidate_plan = base.plan.model_copy(
        update={
            "routes": routes,
            "total_distance_m": sum(item.total_distance_m for item in routes),
            "total_driving_time_s": sum(item.total_duration_s for item in routes),
        }
    )
    validation = validate_plan(dataset_record.dataset, candidate_plan, base.matrix)
    if not validation.valid:
        return None, "這次換車未通過服務區域、載重或時段驗證。", validation, source, target
    return candidate_plan, None, validation, source, target


def _cross_vehicle_route_order_preview_payload(
    base: PlanRecord,
    dataset_record: DatasetRecord,
    payload: CrossVehicleRouteOrderRequest,
) -> dict[str, Any]:
    candidate_plan, reason, validation, source_before, target_before = (
        _cross_vehicle_route_order_candidate(base, dataset_record, payload)
    )
    source_after = next(
        (item for item in (candidate_plan.routes if candidate_plan else [])
         if item.vehicle_id == payload.source_vehicle_id),
        source_before,
    )
    target_after = next(
        (item for item in (candidate_plan.routes if candidate_plan else [])
         if item.vehicle_id == payload.target_vehicle_id),
        target_before,
    )
    before_routes = [route for route in (source_before, target_before) if route is not None]
    after_routes = [route for route in (source_after, target_after) if route is not None]
    before_total_distance = sum(route.total_distance_m for route in before_routes)
    after_total_distance = sum(route.total_distance_m for route in after_routes)
    before_total_duration = sum(route.total_duration_s for route in before_routes)
    after_total_duration = sum(route.total_duration_s for route in after_routes)
    eta_changes: list[dict[str, Any]] = []
    before_etas = {
        stop.order_id: stop.eta
        for route in before_routes
        for stop in route.stops
    }
    after_etas = {
        stop.order_id: stop.eta
        for route in after_routes
        for stop in route.stops
    }
    for order_id, before_eta in before_etas.items():
        after_eta = after_etas.get(order_id)
        if after_eta is None:
            continue
        delta = round(
            (
                datetime.fromisoformat(after_eta)
                - datetime.fromisoformat(before_eta)
            ).total_seconds()
            / 60,
            1,
        )
        if delta:
            eta_changes.append(
                {
                    "order_id": order_id,
                    "before_eta": before_eta,
                    "after_eta": after_eta,
                    "delta_minutes": delta,
                }
            )
    return {
        "plan_id": base.plan_id,
        "base_version": base.version,
        "preview_version": max(store.plans.get(base.plan_id, {0: None})) + 1,
        "source_vehicle_id": payload.source_vehicle_id,
        "target_vehicle_id": payload.target_vehicle_id,
        "order_id": payload.order_id,
        "target_sequence": payload.target_sequence,
        "feasible": candidate_plan is not None and reason is None and validation.valid,
        "reason": reason,
        "requires_human_confirmation": True,
        "source_before": _route_snapshot(source_before) if source_before else None,
        "source_after": _route_snapshot(source_after) if source_after else None,
        "target_before": _route_snapshot(target_before) if target_before else None,
        "target_after": _route_snapshot(target_after) if target_after else None,
        "diff": {
            "distance_delta_m": after_total_distance - before_total_distance,
            "duration_delta_s": after_total_duration - before_total_duration,
            "source_load_delta_kg": (source_after.planned_load_kg - source_before.planned_load_kg)
            if source_after and source_before else 0,
            "target_load_delta_kg": (target_after.planned_load_kg - target_before.planned_load_kg)
            if target_after and target_before else 0,
            "eta_changes": eta_changes,
        },
        "validator": validation.model_dump(mode="json"),
        "provider_mode": base.matrix.provider_mode,
        "matrix_hash": matrix_hash(base.matrix),
    }


def _route_order_preview_payload(
    base: PlanRecord,
    dataset_record: DatasetRecord,
    payload: RouteOrderRequest,
) -> dict[str, Any]:
    before_route, candidate_plan, reason, validation = _route_order_candidate(
        base, dataset_record, payload
    )
    after_route = (
        next(
            (item for item in candidate_plan.routes if item.vehicle_id == payload.vehicle_id),
            before_route,
        )
        if candidate_plan is not None
        else before_route
    )
    before = _route_snapshot(before_route)
    after = _route_snapshot(after_route)
    before_etas = {item["order_id"]: item["eta"] for item in before["stops"]}
    after_etas = {item["order_id"]: item["eta"] for item in after["stops"]}
    eta_changes: list[dict[str, Any]] = []
    for order_id in payload.order_ids:
        before_eta = before_etas.get(order_id)
        after_eta = after_etas.get(order_id)
        if before_eta is None or after_eta is None:
            continue
        before_time = datetime.fromisoformat(before_eta)
        after_time = datetime.fromisoformat(after_eta)
        delta_minutes = round((after_time - before_time).total_seconds() / 60, 1)
        if delta_minutes != 0:
            eta_changes.append(
                {
                    "order_id": order_id,
                    "before_eta": before_eta,
                    "after_eta": after_eta,
                    "delta_minutes": delta_minutes,
                }
            )
    feasible = reason is None and candidate_plan is not None and validation.valid
    preview_version = max(store.plans.get(base.plan_id, {0: None})) + 1
    return {
        "plan_id": base.plan_id,
        "base_version": base.version,
        "preview_version": preview_version,
        "vehicle_id": payload.vehicle_id,
        "order_ids": list(payload.order_ids),
        "feasible": feasible,
        "reason": reason,
        "requires_human_confirmation": True,
        "before": before,
        "after": after,
        "diff": {
            "distance_delta_m": after["total_distance_m"] - before["total_distance_m"],
            "duration_delta_s": after["total_duration_s"] - before["total_duration_s"],
            "load_delta_kg": after["planned_load_kg"] - before["planned_load_kg"],
            "load_utilization_delta": after["load_utilization"] - before["load_utilization"],
            "eta_changes": eta_changes,
        },
        "validator": validation.model_dump(mode="json"),
        "provider_mode": base.matrix.provider_mode,
        "matrix_hash": matrix_hash(base.matrix),
    }


def _dispatch_rule_payload(rule: DispatchRule, *, active_now: bool | None = None) -> dict[str, Any]:
    payload = rule.model_dump(mode="json")
    payload["summary"] = rule_summary(rule)
    payload["active_now"] = rule.active if active_now is None else active_now
    return payload


def _deserialize_matrix(payload: dict[str, Any]) -> MatrixResult:
    return MatrixResult(
        node_ids=tuple(payload["node_ids"]),
        distance_m=tuple(tuple(row) for row in payload["distance_m"]),
        duration_s=tuple(tuple(row) for row in payload["duration_s"]),
        provider_mode=payload.get("provider_mode", "SIMULATED"),
        matrix_version=payload.get("matrix_version", "sim-v1"),
        warning=payload.get("warning"),
    )


def _readable_unassigned_reasons(
    dataset: Dataset | None, plan: PlanResult
) -> dict[str, str]:
    """Classify solver omissions with deterministic, operator-facing reasons."""
    reasons = dict(plan.unassigned_reasons)
    if dataset is None:
        return reasons
    orders = {order.order_id: order for order in dataset.orders}
    routes = {route.vehicle_id: route for route in plan.routes}
    vehicles = {vehicle.vehicle_id: vehicle for vehicle in dataset.vehicles}
    for order_id, reason in reasons.items():
        if reason not in {"UNASSIGNABLE", "UNASSIGNED_BY_SOLVER"}:
            continue
        order = orders.get(order_id)
        if order is None:
            continue
        eligible = [
            vehicle
            for vehicle in vehicles.values()
            if vehicle.status == VehicleStatus.AVAILABLE
            and order.zone_code in vehicle.service_zone_codes
        ]
        if not eligible:
            reasons[order_id] = "SERVICE_ZONE_UNAVAILABLE"
            continue
        # 「每一台車的餘裕都不夠」跟「這張單比最大的那台車還重」是兩件事。
        # ORD-050 是 170 公斤, 最大的車上限 160 公斤, 空車也裝不下; 講成前者
        # 會讓人去看車子還很空, 然後覺得系統在亂講。
        if all(order.total_weight_kg > vehicle.max_load_kg for vehicle in eligible):
            reasons[order_id] = "OVER_VEHICLE_CAPACITY"
            continue
        over_capacity = True
        for vehicle in eligible:
            route = routes.get(vehicle.vehicle_id)
            current_load = route.planned_load_kg if route is not None else vehicle.current_load_kg
            if current_load + order.total_weight_kg <= vehicle.max_load_kg:
                over_capacity = False
                break
        reasons[order_id] = "CAPACITY_LIMIT" if over_capacity else "TIME_OR_ROUTE_CONFLICT"
    return reasons


def _deserialize_dataset(payload: dict[str, Any]) -> Dataset:
    package_values = [Package.model_validate(package) for package in payload["packages"]]
    package_by_order: dict[str, list[Package]] = {}
    for package in package_values:
        package_by_order.setdefault(package.order_id, []).append(package)
    orders: list[Order] = []
    for raw_order in payload["orders"]:
        order_data = dict(raw_order)
        order_data.pop("total_weight_kg", None)
        order_data["priority"] = Priority(order_data.get("priority", Priority.NORMAL.value))
        order_data["packages"] = tuple(package_by_order.get(order_data["order_id"], ()))
        orders.append(Order.model_validate(order_data))
    vehicles = []
    for raw_vehicle in payload["vehicles"]:
        vehicle_data = dict(raw_vehicle)
        vehicle_data["status"] = VehicleStatus(vehicle_data["status"])
        vehicle_data["service_zone_codes"] = tuple(vehicle_data.get("service_zone_codes", ()))
        vehicles.append(Vehicle.model_validate(vehicle_data))
    zones = []
    for raw_zone in payload["zones"]:
        zone_data = dict(raw_zone)
        for field_name in (
            "covered_cities",
            "covered_districts",
            "tdx_city_codes",
            "adjacent_zone_codes",
        ):
            zone_data[field_name] = tuple(zone_data.get(field_name, ()))
        zones.append(Zone.model_validate(zone_data))
    return Dataset.model_validate(
        {
            "orders": tuple(orders),
            "packages": tuple(package_values),
            "vehicles": tuple(vehicles),
            "zones": tuple(zones),
            "source_filename": payload.get("source_filename", "workbook.xlsx"),
        }
    )


def _hydrate_store() -> None:
    """Restore non-secret metadata and immutable versions after a local restart."""
    for row in repository.load_datasets():
        try:
            dataset = _deserialize_dataset(json.loads(row["payload_json"]))
            dataset_validation = ValidationReport.model_validate(json.loads(row["validation_json"]))
            matrix = _deserialize_matrix(json.loads(row["matrix_json"]))
            store.add_dataset(
                DatasetRecord(
                    dataset_id=row["dataset_id"],
                    dataset=dataset,
                    validation=dataset_validation,
                    matrix=matrix,
                    created_at=row["created_at"],
                )
            )
        except (KeyError, TypeError, ValueError):
            continue
    for row in repository.load_plans():
        try:
            plan = PlanResult.model_validate(json.loads(row["payload_json"]))
            plan_validation = PlanValidation.model_validate(json.loads(row["validation_json"]))
            matrix = _deserialize_matrix(json.loads(row["matrix_json"]))
            store.add_plan(
                PlanRecord(
                    plan_id=row["plan_id"],
                    dataset_id=row["dataset_id"],
                    version=int(row["version"]),
                    state=row["state"],
                    plan=plan,
                    validation=plan_validation,
                    matrix=matrix,
                    created_at=row["created_at"],
                ),
                make_current=False,
            )
        except (KeyError, TypeError, ValueError):
            continue
    store.current_versions.update(repository.current_versions())


_hydrate_store()


@app.middleware("http")
async def request_id_middleware(request: Request, call_next: Any) -> Any:
    request_id = request.headers.get("X-Request-ID") or f"REQ-{uuid4().hex[:12]}"
    request.state.request_id = request_id
    response = await call_next(request)
    response.headers["X-Request-ID"] = request_id
    return response


@app.get("/health")
def health(request: Request) -> dict[str, str]:
    return {
        "status": "ok",
        "service": "ai-delivery-dispatch-agent",
        "request_id": _request_id(request),
    }


@app.get("/ready")
def ready(request: Request) -> dict[str, Any]:
    credential_status = settings.credential_status()
    return {
        "status": "ready",
        "components": {
            "database": "ready",
            "optimizer": "ready",
            "openai": "ready"
            if credential_status["OPENAI_API_KEY"] == "CONFIGURED"
            else "degraded",
            "google_routes": "enabled"
            if settings.google_routes_enabled
            and credential_status["GOOGLE_ROUTES_SERVER_API_KEY"] == "CONFIGURED"
            else "disabled",
            "tdx": "enabled"
            if credential_status["TDX_CLIENT_ID"] == "CONFIGURED"
            and credential_status["TDX_CLIENT_SECRET"] == "CONFIGURED"
            else "disabled",
        },
        "request_id": _request_id(request),
    }


@app.get("/auth/status", include_in_schema=False)
def auth_status(request: Request) -> dict[str, Any]:
    """Return only the demo gate state; never disclose the configured password."""
    return {
        "required": bool(settings.demo_access_password),
        "authenticated": _has_demo_session(request) if settings.demo_access_password else True,
    }


@app.post("/auth/login", include_in_schema=False)
def auth_login(payload: DemoLoginRequest, request: Request) -> Any:
    expected = settings.demo_access_password
    if not expected:
        return {"authenticated": True, "required": False}
    if not hmac.compare_digest(payload.password, expected):
        return _error(request, 401, "DEMO_AUTH_INVALID", "展示環境密碼不正確。")
    response = JSONResponse(content={"authenticated": True, "required": True})
    response.set_cookie(
        "dispatch_demo_session",
        _demo_session_token() or "",
        max_age=60 * 60 * 12,
        httponly=True,
        secure=request.url.scheme == "https",
        samesite="lax",
    )
    return response


@app.post("/auth/logout", include_in_schema=False)
def auth_logout() -> JSONResponse:
    response = JSONResponse(content={"authenticated": False})
    response.delete_cookie("dispatch_demo_session")
    return response


@app.get("/api/v1/runtime-config", include_in_schema=False)
def runtime_config() -> dict[str, str | None]:
    """Expose only the browser-safe runtime configuration to the SPA."""
    return {
        "google_maps_browser_api_key": settings.google_maps_browser_api_key
        if settings.google_routes_enabled
        else None
    }


def _workbook_headers_and_samples(
    content: bytes,
) -> tuple[dict[str, list[str]], dict[str, list[list[str]]]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        headers: dict[str, list[str]] = {}
        samples: dict[str, list[list[str]]] = {}
        for sheet_name in workbook.sheetnames:
            rows = list(workbook[sheet_name].iter_rows(min_row=1, max_row=4, values_only=True))
            headers[sheet_name] = [
                str(value).strip() if value is not None else ""
                for value in (rows[0] if rows else ())
            ]
            samples[sheet_name] = [
                ["" if value is None else str(value)[:80] for value in row]
                for row in rows[1:]
            ]
        return headers, samples
    finally:
        workbook.close()


def _mapping_response(
    source_name: str,
    headers_by_sheet: dict[str, list[str]],
    samples_by_sheet: dict[str, list[list[str]]],
    mapping: ColumnMapping,
    *,
    status: str,
    requires_confirmation: bool,
    confidence_by_source: dict[tuple[str, str], float] | None = None,
) -> dict[str, Any]:
    required_fields = {
        "orders": {"location_label", "time_slot"},
        "packages": {"weight_kg"},
    }
    entries: list[dict[str, Any]] = []
    missing_fields: list[str] = []
    for sheet_name, headers in headers_by_sheet.items():
        sheet_mapping = mapping.get(sheet_name, {})
        targets = set(sheet_mapping.values())
        for required_field in sorted(required_fields.get(sheet_name, set()) - targets):
            missing_fields.append(f"{sheet_name}.{required_field}")
        for index, source in enumerate(headers):
            target = sheet_mapping.get(source)
            confidence = (
                confidence_by_source.get((sheet_name, source), 0.0)
                if confidence_by_source is not None
                else 1.0
            )
            sample_values = [
                row[index] for row in samples_by_sheet.get(sheet_name, []) if index < len(row)
            ]
            entries.append(
                {
                    "sheet": sheet_name,
                    "source": source,
                    "target": target,
                    "confidence": confidence,
                    "sample_values": sample_values[:3],
                    "required": target in required_fields.get(sheet_name, set()),
                }
            )
    return {
        "status": status,
        "source_name": source_name,
        "requires_confirmation": requires_confirmation,
        "entries": entries,
        "missing_fields": missing_fields,
        "mapping": mapping,
    }


def _inspection_error(
    request: Request,
    source_name: str,
    status_code: int,
    code: str,
    message: str,
    field_errors: list[dict[str, Any]] | None = None,
) -> Any:
    """Keep browser preflight failures in-band while preserving REST error codes."""
    if request.headers.get("X-Dispatch-UI") == "true":
        return {
            "status": "INVALID",
            "source_name": source_name,
            "requires_confirmation": False,
            "entries": [],
            "missing_fields": [],
            "mapping": {},
            "error": {
                "code": code,
                "message": message,
                "field_errors": field_errors or [],
            },
            "request_id": _request_id(request),
        }
    return _error(
        request,
        status_code,
        code,
        message,
        field_errors=field_errors,
    )


@app.post("/api/v1/datasets/inspect-excel")
async def inspect_excel(request: Request, file: Annotated[UploadFile, File(...)]) -> Any:
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        return _inspection_error(
            request,
            file.filename or "upload.xlsx",
            400,
            "INVALID_FILE_TYPE",
            "只接受 .xlsx 檔案。",
        )
    content = await file.read()
    if not content:
        return _inspection_error(
            request,
            file.filename,
            422,
            "EMPTY_WORKBOOK",
            "Excel 沒有資料，請提供至少一列訂單。",
        )
    try:
        headers_by_sheet, samples_by_sheet = _workbook_headers_and_samples(content)
    except Exception:
        return _inspection_error(
            request,
            file.filename,
            400,
            "INVALID_XLSX",
            "無法讀取 .xlsx 檔案。",
        )
    if tuple(headers_by_sheet) != tuple(SHEET_FIELDS):
        return _inspection_error(
            request,
            file.filename,
            422,
            "INVALID_SHEETS",
            "必須包含且只包含 orders、packages、vehicles、zones 四張工作表。",
        )
    canonical = all(headers_by_sheet[name] == list(SHEET_FIELDS[name]) for name in SHEET_FIELDS)
    saved = saved_mapping_profiles.get(file.filename)
    if canonical:
        payload = _mapping_response(
            file.filename,
            headers_by_sheet,
            samples_by_sheet,
            {},
            status="CANONICAL",
            requires_confirmation=False,
        )
    elif saved is not None:
        payload = _mapping_response(
            file.filename,
            headers_by_sheet,
            samples_by_sheet,
            saved,
            status="AUTO_APPLIED",
            requires_confirmation=False,
        )
    else:
        try:
            proposal = await propose_mapping(headers_by_sheet, SHEET_FIELDS)
        except Exception as exc:
            status_code, error_code, message, retryable = _classify_agent_error(exc)
            if status_code not in {503, 504}:
                return _error(
                    request,
                    status_code,
                    error_code,
                    message,
                    provider="OPENAI",
                    exception_type=type(exc).__name__,
                    retryable=retryable,
                )
            provider_runtime_state["openai"] = "failed"
            fallback = positional_mapping(headers_by_sheet, SHEET_FIELDS)
            fallback_mapping = {
                name: {
                    item.source: item.target
                    for item in fallback.suggestions
                    if item.sheet == name and item.target is not None
                }
                for name in SHEET_FIELDS
            }
            fallback_confidence_by_source = {
                (item.sheet, item.source): item.confidence for item in fallback.suggestions
            }
            payload = _mapping_response(
                file.filename,
                headers_by_sheet,
                samples_by_sheet,
                fallback_mapping,
                status="NEEDS_CONFIRMATION",
                requires_confirmation=True,
                confidence_by_source=fallback_confidence_by_source,
            )
            payload["provider_fallback"] = {
                "provider": "OPENAI",
                "error_code": error_code,
                "message": "AI 對映服務目前無法連線，已提供低信心度結構草稿，必須人工確認。",
            }
        else:
            allowed = {name: set(fields) for name, fields in SHEET_FIELDS.items()}
            mapping: ColumnMapping = {name: {} for name in SHEET_FIELDS}
            confidence_by_source: dict[tuple[str, str], float] = {}
            for suggestion in proposal.suggestions:
                if suggestion.sheet not in allowed:
                    continue
                if suggestion.source not in headers_by_sheet.get(suggestion.sheet, []):
                    continue
                if (
                    suggestion.target is not None
                    and suggestion.target not in allowed[suggestion.sheet]
                ):
                    continue
                if suggestion.target is not None:
                    mapping[suggestion.sheet][suggestion.source] = suggestion.target
                confidence_by_source[(suggestion.sheet, suggestion.source)] = suggestion.confidence
            payload = _mapping_response(
                file.filename,
                headers_by_sheet,
                samples_by_sheet,
                mapping,
                status="NEEDS_CONFIRMATION",
                requires_confirmation=True,
                confidence_by_source=confidence_by_source,
            )
    if not payload["requires_confirmation"]:
        _, report = parse_workbook(
            BytesIO(content), source_filename=file.filename, column_mapping=saved
        )
        if not report.is_valid:
            return _inspection_error(
                request,
                file.filename,
                422,
                "DATASET_VALIDATION_FAILED",
                "工作簿驗證失敗。",
                field_errors=[error.model_dump() for error in report.errors],
            )
    payload["request_id"] = _request_id(request)
    return payload


@app.post("/api/v1/datasets/import-excel", status_code=201)
async def import_excel(
    request: Request,
    file: Annotated[UploadFile, File(...)],
    mapping: Annotated[str | None, Form()] = None,
    mapping_name: Annotated[str | None, Form()] = None,
) -> Any:
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        return _error(request, 400, "INVALID_FILE_TYPE", "只接受 .xlsx 檔案。")
    content = await file.read()
    if len(content) > 20 * 1024 * 1024:
        return _error(request, 413, "FILE_TOO_LARGE", "工作簿超過大小限制。")
    selected_mapping = saved_mapping_profiles.get(file.filename)
    if mapping:
        try:
            mapping_payload = ColumnMappingRequest.model_validate(
                {"mapping": json.loads(mapping), "save_as": mapping_name}
            )
            selected_mapping = mapping_payload.mapping
        except Exception:
            return _error(
                request,
                422,
                "INVALID_COLUMN_MAPPING",
                "欄位對映格式不正確，請重新確認。",
            )
    dataset, report = parse_workbook(
        BytesIO(content), source_filename=file.filename, column_mapping=selected_mapping
    )
    if dataset is None or not report.is_valid:
        return _error(
            request,
            422,
            "DATASET_VALIDATION_FAILED",
            "工作簿驗證失敗。",
            field_errors=[error.model_dump() for error in report.errors],
            requires_manual_review=report.requires_manual_review,
        )
    dataset_id = f"DS-{uuid4().hex[:12].upper()}"
    record = DatasetRecord(
        dataset_id=dataset_id,
        dataset=dataset,
        validation=report,
        matrix=SimulatedRouteProvider().build(dataset),
        created_at=datetime.now(UTC).isoformat(),
    )
    store.add_dataset(record)
    repository.save_dataset(dataset_id, dataset, report, record.matrix, record.created_at)
    if mapping and mapping_name:
        saved_mapping_profiles[file.filename] = selected_mapping or {}
        saved_mapping_profiles[mapping_name] = selected_mapping or {}
    return {
        "dataset_id": dataset_id,
        "status": "VALIDATED",
        "counts": {
            "orders": len(dataset.orders),
            "packages": len(dataset.packages),
            "vehicles": len(dataset.vehicles),
            "zones": len(dataset.zones),
        },
        "total_weight_kg": round(sum(order.total_weight_kg for order in dataset.orders), 3),
        "validation": _validation_payload(report),
        "request_id": _request_id(request),
    }


@app.get("/api/v1/datasets/{dataset_id}")
def get_dataset(dataset_id: str, request: Request) -> Any:
    record = store.get_dataset(dataset_id)
    if record is None:
        return _error(request, 404, "DATASET_NOT_FOUND", "找不到資料集。")
    dataset = record.dataset
    return {
        "dataset_id": dataset_id,
        "source_filename": dataset.source_filename,
        "status": "VALIDATED" if record.validation.is_valid else "INVALID",
        "counts": {
            "orders": len(dataset.orders),
            "packages": len(dataset.packages),
            "vehicles": len(dataset.vehicles),
            "zones": len(dataset.zones),
        },
        "total_weight_kg": round(sum(order.total_weight_kg for order in dataset.orders), 3),
        "created_at": record.created_at,
        "request_id": _request_id(request),
    }


@app.get("/api/v1/datasets/{dataset_id}/validation")
def get_dataset_validation(dataset_id: str, request: Request) -> Any:
    record = store.get_dataset(dataset_id)
    if record is None:
        return _error(request, 404, "DATASET_NOT_FOUND", "找不到資料集。")
    return {
        "dataset_id": dataset_id,
        "validation": _validation_payload(record.validation),
        "request_id": _request_id(request),
    }


@app.post("/api/v1/plans", status_code=201)
def create_plan(payload: CreatePlanRequest, request: Request) -> Any:
    dataset_record = store.get_dataset(payload.dataset_id)
    if dataset_record is None:
        return _error(request, 404, "DATASET_NOT_FOUND", "找不到資料集。")
    prefer_live = payload.route_provider_preference == "AUTO" and payload.traffic_mode == "AUTO"
    try:
        matrix = _build_matrix(dataset_record.dataset, prefer_live=prefer_live)
    except GoogleRoutesProviderError as exc:
        return _error(
            request,
            502,
            "PROVIDER_UNAVAILABLE",
            "Google Routes 即時矩陣無法取得。",
            provider="GOOGLE",
            operation="computeRouteMatrix",
            provider_error=exc.code,
            provider_error_category=exc.category,
            fallback_used=False,
            retryable=exc.code in {"GOOGLE_TIMEOUT", "GOOGLE_REQUEST_FAILED"},
        )
    if payload.algorithm == "BASELINE":
        plan = build_baseline(dataset_record.dataset, matrix)
    else:
        plan = build_plan_with_rules(
            dataset_record.dataset,
            matrix,
            settings.solver_time_limit_seconds,
            objective=payload.objective,
            rules=list_dispatch_rules(include_inactive=False),
        )
        plan = apply_service_time_parameters(plan, dataset_record.dataset, matrix)
    validation = validate_plan(dataset_record.dataset, plan, matrix)
    plan_id = f"PLAN-{uuid4().hex[:12].upper()}"
    record = PlanRecord(
        plan_id=plan_id,
        dataset_id=payload.dataset_id,
        version=1,
        state="PROPOSED",
        plan=plan,
        validation=validation,
        matrix=matrix,
        created_at=datetime.now(UTC).isoformat(),
    )
    store.add_plan(record)
    repository.save_plan(record)
    if not validation.valid:
        return _error(
            request, 409, "PLAN_NOT_CONFIRMABLE", "規劃結果未通過獨立驗證。", plan_id=plan_id
        )
    return _plan_payload(record) | {"request_id": _request_id(request)}


@app.post("/api/v1/plans/compare")
def compare_plan_strategies(payload: CompareStrategiesRequest, request: Request) -> Any:
    """Solve the three supported objectives against one shared matrix."""
    dataset_record = store.get_dataset(payload.dataset_id)
    if dataset_record is None:
        return _error(request, 404, "DATASET_NOT_FOUND", "找不到資料集。")
    if payload.plan_id is not None:
        plan_record = store.get_plan(payload.plan_id, payload.version)
        if plan_record is None:
            return _error(request, 404, "PLAN_NOT_FOUND", "找不到策略比較使用的方案版本。")
        if plan_record.dataset_id != payload.dataset_id:
            return _error(request, 409, "PLAN_DATASET_MISMATCH", "方案與資料集不一致。")
        matrix = plan_record.matrix
    else:
        prefer_live = (
            payload.route_provider_preference == "AUTO" and payload.traffic_mode == "AUTO"
        )
        try:
            matrix = _build_matrix(dataset_record.dataset, prefer_live=prefer_live)
        except GoogleRoutesProviderError as exc:
            return _error(
                request,
                502,
                "PROVIDER_UNAVAILABLE",
                "Google Routes 即時矩陣無法取得。",
                provider="GOOGLE",
                provider_error=exc.code,
                provider_error_category=exc.category,
                fallback_used=False,
            )
    strategy_goals: dict[str, tuple[str, str]] = {
        "FASTEST": (
            "最小化總行駛時間（秒）",
            "速度優先，可能犧牲車輛工作量平衡。",
        ),
        "BALANCED": (
            "縮小各車載重差距",
            "工作量較平均，可能增加總距離與行駛時間。",
        ),
        "STABLE": (
            "保留較大的時段餘裕",
            "較能承受延遲，可能增加總距離與時間。",
        ),
    }
    strategies: list[dict[str, Any]] = []
    for objective in ("FASTEST", "BALANCED", "STABLE"):
        plan = build_ortools(
            dataset_record.dataset,
            matrix,
            settings.solver_time_limit_seconds,
            objective=objective,
        )
        validation = validate_plan(dataset_record.dataset, plan, matrix)
        loads = [route.planned_load_kg for route in plan.routes]
        risks = calculate_plan_risks(dataset_record.dataset, plan)
        min_slack = min((risk["slack_minutes"] for risk in risks), default=0.0)
        strategies.append(
            {
                "objective": objective,
                "primary_goal": strategy_goals[objective][0],
                "tradeoff": strategy_goals[objective][1],
                "algorithm": plan.algorithm,
                "total_distance_m": plan.total_distance_m,
                "total_duration_s": plan.total_driving_time_s,
                "max_vehicle_load_kg": max(loads, default=0.0),
                "load_spread_kg": round(max(loads, default=0.0) - min(loads, default=0.0), 3),
                "min_slack_minutes": min_slack,
                "unassigned_orders": plan.unassigned_orders,
                "validator": validation.model_dump(mode="json"),
            }
        )
    return {
        "dataset_id": payload.dataset_id,
        "dataset_hash": dataset_hash(dataset_record.dataset),
        "matrix_hash": matrix_hash(matrix),
        "matrix_version": matrix.matrix_version,
        "provider_mode": matrix.provider_mode,
        "strategies": strategies,
        "request_id": _request_id(request),
    }


@app.get("/api/v1/dispatch-rules")
def get_dispatch_rules(request: Request) -> Any:
    rules = list_dispatch_rules(include_inactive=True)
    active_rules = {rule.rule_id for rule in list_dispatch_rules(include_inactive=False)}
    return {
        "rules": [
            _dispatch_rule_payload(rule, active_now=rule.rule_id in active_rules)
            for rule in rules
        ],
        "active_count": len(active_rules),
        "request_id": _request_id(request),
    }


@app.post("/api/v1/dispatch-rules/confirm", status_code=201)
def confirm_dispatch_rule(payload: DispatchRuleMutationRequest, request: Request) -> Any:
    base_record = store.get_plan(payload.plan_id, payload.base_plan_version)
    if base_record is None:
        return _error(request, 404, "PLAN_NOT_FOUND", "找不到規則試算所需的規劃版本。")
    dataset_record = store.get_dataset(base_record.dataset_id)
    if dataset_record is None:
        return _error(request, 404, "DATASET_NOT_FOUND", "找不到規則試算所需的資料集。")
    draft = DispatchRuleDraft(
        subject_type=payload.subject_type,
        subject_id=payload.subject_id,
        rule_type=payload.rule_type,
        value=payload.value,
        duration=payload.duration,
    )
    trial = preview_rule_trial(
        dataset_record.dataset,
        base_record.matrix,
        base_record.plan,
        draft,
        payload.source_utterance,
        settings.solver_time_limit_seconds,
        list_dispatch_rules(include_inactive=False),
    )
    additional_rule: DispatchRule | None = None
    if payload.additional_rule is not None:
        primary_rule = DispatchRule(
            rule_id="RULE-CANDIDATE-PRIMARY",
            subject_type=payload.subject_type,
            subject_id=payload.subject_id,
            rule_type=payload.rule_type,
            value=payload.value,
            source_utterance=payload.source_utterance,
            created_at="1970-01-01T00:00:00+00:00",
            expires_at=expires_at_for_duration(payload.duration),
        )
        additional_rule = DispatchRule(
            rule_id="RULE-CANDIDATE-SECONDARY",
            subject_type=payload.subject_type,
            subject_id=payload.subject_id,
            rule_type=payload.additional_rule.rule_type,
            value=payload.additional_rule.value,
            source_utterance=payload.source_utterance,
            created_at="1970-01-01T00:00:00+00:00",
            expires_at=expires_at_for_duration(payload.additional_rule.duration),
        )
        combined_trial = preview_rule_trial(
            dataset_record.dataset,
            base_record.matrix,
            base_record.plan,
            DispatchRuleDraft(
                subject_type=payload.subject_type,
                subject_id=payload.subject_id,
                rule_type=additional_rule.rule_type,
                value=additional_rule.value,
                duration=payload.additional_rule.duration,
            ),
            payload.source_utterance,
            settings.solver_time_limit_seconds,
            [*list_dispatch_rules(include_inactive=False), primary_rule],
        )
        trial = combined_trial.model_copy(
            update={
                "status": (
                    "FEASIBLE"
                    if trial.status == "FEASIBLE"
                    and combined_trial.status == "FEASIBLE"
                    else "CONFLICT"
                ),
                "affected_order_ids": sorted(
                    set(trial.affected_order_ids)
                    | set(combined_trial.affected_order_ids)
                ),
                "conflicts": [*trial.conflicts, *combined_trial.conflicts],
            }
        )
    if trial.status != "FEASIBLE":
        return _error(
            request,
            409,
            "DISPATCH_RULE_CONFLICT",
            "這條規則試算未通過，尚未建立。",
            conflicts=[item.model_dump(mode="json") for item in trial.conflicts],
            affected_order_ids=trial.affected_order_ids,
            trial=trial.model_dump(mode="json", exclude={"plan"}),
        )
    rule = DispatchRule(
        rule_id=f"RULE-{uuid4().hex[:12].upper()}",
        subject_type=payload.subject_type,
        subject_id=payload.subject_id,
        rule_type=payload.rule_type,
        value=payload.value,
        source_utterance=payload.source_utterance,
        created_at=datetime.now(UTC).isoformat(),
        expires_at=expires_at_for_duration(payload.duration),
    )
    save_dispatch_rule(rule)
    if additional_rule is not None:
        save_dispatch_rule(
            additional_rule.model_copy(
                update={"rule_id": f"RULE-{uuid4().hex[:12].upper()}"}
            )
        )
    repository.append_audit(
        event_id=f"AUD-{uuid4().hex[:12].upper()}",
        event_type="DISPATCH_RULE_CREATED",
        created_at=datetime.now(UTC).isoformat(),
        plan_id=base_record.plan_id,
        version=base_record.version,
        metadata={"rule": rule.model_dump(mode="json")},
    )
    active_count = len(list_dispatch_rules(include_inactive=False))
    return {
        "rule": _dispatch_rule_payload(rule),
        "active_count": active_count,
        "trial": trial.model_dump(mode="json", exclude={"plan"}),
        "request_id": _request_id(request),
    }


@app.post("/api/v1/dispatch-rules/{rule_id}/deactivate")
def deactivate_rule(rule_id: str, request: Request) -> Any:
    rule = deactivate_dispatch_rule(rule_id)
    if rule is None:
        return _error(request, 404, "DISPATCH_RULE_NOT_FOUND", "找不到要停用的規則。")
    repository.append_audit(
        event_id=f"AUD-{uuid4().hex[:12].upper()}",
        event_type="DISPATCH_RULE_DEACTIVATED",
        created_at=datetime.now(UTC).isoformat(),
        metadata={"rule_id": rule_id},
    )
    return {
        "rule": _dispatch_rule_payload(rule, active_now=False),
        "active_count": len(list_dispatch_rules(include_inactive=False)),
        "request_id": _request_id(request),
    }


@app.get("/api/v1/dispatch-parameters")
def get_dispatch_parameters(request: Request) -> Any:
    return parameter_state() | {"request_id": _request_id(request)}


@app.post("/api/v1/runtime/reset")
def reset_runtime_state(request: Request) -> Any:
    """Reset persistent demo-only parameters and dispatch rules."""
    cleared_rule_count = reset_dispatch_rules()
    parameters = reset_service_parameters()
    return {
        **parameters,
        "cleared_rule_count": cleared_rule_count,
        "request_id": _request_id(request),
    }


@app.post("/api/v1/dispatch-parameters/confirm", status_code=201)
def confirm_dispatch_parameter(
    payload: DispatchParameterConfirmationRequest, request: Request
) -> Any:
    base_record = store.get_plan(payload.plan_id, payload.base_plan_version)
    if base_record is None:
        return _error(request, 404, "PLAN_NOT_FOUND", "找不到參數確認所需的規劃版本。")
    if stage_for_state(base_record.state) != "DISPATCHED":
        return _error(
            request,
            409,
            "DISPATCH_PARAMETER_STAGE_INVALID",
            "配送偏差參數只能在已發車的時間軸回顧中確認。",
        )
    try:
        confirmed = confirm_service_parameter(
            payload.zone_code,
            payload.from_service_minutes,
            payload.to_service_minutes,
        )
    except ValueError as exc:
        error_code = str(exc)
        message = (
            "參數已被其他操作更新，請重新讀取目前值。"
            if error_code == "PARAMETER_VERSION_CONFLICT"
            else "服務時間參數不在允許範圍內。"
        )
        return _error(request, 409, error_code, message)
    repository.append_audit(
        event_id=f"AUD-{uuid4().hex[:12].upper()}",
        event_type="DISPATCH_PARAMETER_CONFIRMED",
        created_at=datetime.now(UTC).isoformat(),
        plan_id=base_record.plan_id,
        version=base_record.version,
        metadata={"source": payload.source, **confirmed},
    )
    return {**confirmed, "request_id": _request_id(request)}


@app.get("/api/v1/plans/{plan_id}")
def get_plan(plan_id: str, request: Request, version: int | None = None) -> Any:
    record = store.get_plan(plan_id, version)
    if record is None:
        return _error(request, 404, "PLAN_NOT_FOUND", "找不到規劃版本。")
    return _plan_payload(record) | {"request_id": _request_id(request)}


@app.get("/api/v1/plans/{plan_id}/versions")
def list_plan_versions(plan_id: str, request: Request) -> Any:
    versions = store.plans.get(plan_id)
    if not versions:
        return _error(request, 404, "PLAN_NOT_FOUND", "找不到規劃版本。")
    current = store.current_versions.get(plan_id)
    return {
        "plan_id": plan_id,
        "current_version": current,
        "versions": [
            {
                "version": record.version,
                "state": record.state,
                "created_at": record.created_at,
                "algorithm": record.plan.algorithm,
                "objective": record.plan.objective,
                "validator_valid": record.validation.valid,
                "complete": record.plan.complete,
                "unassigned_orders": record.plan.unassigned_orders,
            }
            for record in sorted(versions.values(), key=lambda item: item.version)
        ],
        "request_id": _request_id(request),
    }


@app.post("/api/v1/plans/{plan_id}/restore")
def restore_plan(plan_id: str, payload: RestorePlanRequest, request: Request) -> Any:
    source = store.get_plan(plan_id, payload.source_version)
    if source is None:
        return _error(request, 404, "PLAN_NOT_FOUND", "找不到要復原的規劃版本。")
    if source.state == "DISPATCHED":
        return _error(request, 409, "PLAN_ALREADY_DISPATCHED", "已出發的規劃不可復原。")
    dataset_record = store.get_dataset(source.dataset_id)
    if dataset_record is None:
        return _error(request, 404, "DATASET_NOT_FOUND", "找不到來源資料集。")
    validation = validate_plan(dataset_record.dataset, source.plan, source.matrix)
    if not validation.valid or not source.plan.complete or source.plan.unassigned_orders:
        return _error(request, 409, "PLAN_NOT_CONFIRMABLE", "來源版本未通過完整性驗證，無法復原。")
    next_version = max(store.plans.get(plan_id, {0: None})) + 1
    restored_plan = source.plan.model_copy(update={"state": "PROPOSED"})
    restored = PlanRecord(
        plan_id=plan_id,
        dataset_id=source.dataset_id,
        version=next_version,
        state="PROPOSED",
        plan=restored_plan,
        validation=validation,
        matrix=source.matrix,
        created_at=datetime.now(UTC).isoformat(),
    )
    store.add_plan(restored, make_current=False)
    repository.save_plan(restored, make_current=False)
    repository.append_audit(
        f"AUD-{uuid4().hex[:12].upper()}",
        "PLAN_RESTORED_PREVIEW",
        restored.created_at,
        plan_id,
        next_version,
        {
            "source_version": payload.source_version,
            "dispatcher_reference": payload.dispatcher_reference,
        },
    )
    return _plan_payload(restored) | {
        "restored_from_version": payload.source_version,
        "requires_human_confirmation": True,
        "request_id": _request_id(request),
    }


@app.post("/api/v1/plans/{plan_id}/delay-preview")
def delay_preview(plan_id: str, payload: DelaySimulationRequest, request: Request) -> Any:
    record = store.get_plan(plan_id, payload.version)
    if record is None:
        return _error(request, 404, "PLAN_NOT_FOUND", "找不到規劃版本。")
    dataset_record = store.get_dataset(record.dataset_id)
    if dataset_record is None:
        return _error(request, 404, "DATASET_NOT_FOUND", "找不到規劃資料集。")
    risks = calculate_plan_risks(dataset_record.dataset, record.plan)
    return {
        "plan_id": plan_id,
        "version": record.version,
        "risks": risks,
        "simulation": summarize_delay(record.plan, risks, payload.delay_minutes),
        "validator": record.validation.model_dump(mode="json"),
        "request_id": _request_id(request),
    }


@app.post("/api/v1/plans/{plan_id}/reassign/preview")
def reassignment_preview(plan_id: str, payload: ReassignmentRequest, request: Request) -> Any:
    base = store.get_plan(plan_id, payload.base_plan_version)
    if base is None:
        return _error(request, 404, "PLAN_NOT_FOUND", "找不到基準規劃版本。")
    if base.state == "DISPATCHED":
        return _error(request, 409, "PLAN_ALREADY_DISPATCHED", "已出發的規劃不可換車。")
    dataset_record = store.get_dataset(base.dataset_id)
    if dataset_record is None:
        return _error(request, 404, "DATASET_NOT_FOUND", "找不到基準資料集。")
    preview = preview_reassignment(
        base.plan,
        dataset_record.dataset,
        base.matrix,
        payload.order_id,
        payload.target_vehicle_id,
    )
    if preview is None:
        return _error(
            request,
            409,
            "REASSIGNMENT_NOT_FEASIBLE",
            "換車預覽不符合容量、服務區域或時段限制；原方案未變更。",
            plan_id=plan_id,
            order_id=payload.order_id,
            target_vehicle_id=payload.target_vehicle_id,
            requires_manual_review=True,
        )
    validation = validate_plan(dataset_record.dataset, preview, base.matrix)
    if not validation.valid:
        return _error(
            request,
            409,
            "REASSIGNMENT_NOT_FEASIBLE",
            "換車預覽未通過獨立驗證；原方案未變更。",
            validation=validation.model_dump(mode="json"),
            requires_manual_review=True,
        )
    preview_version = max(store.plans.get(plan_id, {0: None})) + 1
    record = PlanRecord(
        plan_id=plan_id,
        dataset_id=base.dataset_id,
        version=preview_version,
        state="PROPOSED",
        plan=preview,
        validation=validation,
        matrix=base.matrix,
        created_at=datetime.now(UTC).isoformat(),
    )
    store.add_plan(record, make_current=False)
    repository.save_plan(record, make_current=False)
    return {
        "plan_id": plan_id,
        "base_version": base.version,
        "preview_version": preview_version,
        "requires_human_confirmation": True,
        "before": _plan_payload(base)["summary"],
        "after": _plan_payload(record)["summary"],
        "diff": compute_plan_diff(base.plan, preview),
        "validator": validation.model_dump(mode="json"),
        "provider_mode": base.matrix.provider_mode,
        "matrix_hash": matrix_hash(base.matrix),
        "request_id": _request_id(request),
    }


@app.post("/api/v1/plans/{plan_id}/route-order/preview")
def route_order_preview(plan_id: str, payload: RouteOrderRequest, request: Request) -> Any:
    base = store.get_plan(plan_id, payload.base_plan_version)
    if base is None:
        return _error(request, 404, "PLAN_NOT_FOUND", "找不到基準規劃版本。")
    dataset_record = store.get_dataset(base.dataset_id)
    if dataset_record is None:
        return _error(request, 404, "DATASET_NOT_FOUND", "找不到規劃資料集。")
    return _route_order_preview_payload(base, dataset_record, payload) | {
        "request_id": _request_id(request),
    }


@app.post("/api/v1/plans/{plan_id}/route-order/confirm")
def route_order_confirm(plan_id: str, payload: RouteOrderRequest, request: Request) -> Any:
    base = store.get_plan(plan_id, payload.base_plan_version)
    if base is None:
        return _error(request, 404, "PLAN_NOT_FOUND", "找不到基準規劃版本。")
    dataset_record = store.get_dataset(base.dataset_id)
    if dataset_record is None:
        return _error(request, 404, "DATASET_NOT_FOUND", "找不到規劃資料集。")
    _before_route, candidate_plan, reason, validation = _route_order_candidate(
        base, dataset_record, payload
    )
    if candidate_plan is None or reason is not None or not validation.valid:
        return _error(
            request,
            409,
            "ROUTE_ORDER_NOT_FEASIBLE",
            reason or "站序未通過獨立驗證；原方案未變更。",
            validation=validation.model_dump(mode="json"),
            requires_manual_review=True,
        )
    next_version = max(store.plans.get(plan_id, {0: None})) + 1
    record = PlanRecord(
        plan_id=plan_id,
        dataset_id=base.dataset_id,
        version=next_version,
        state=base.state,
        plan=candidate_plan,
        validation=validation,
        matrix=base.matrix,
        created_at=datetime.now(UTC).isoformat(),
    )
    store.add_plan(record)
    repository.save_plan(record)
    return _plan_payload(record) | {
        "request_id": _request_id(request),
    }


@app.post("/api/v1/plans/{plan_id}/route-order/cross-vehicle/preview")
def cross_vehicle_route_order_preview(
    plan_id: str, payload: CrossVehicleRouteOrderRequest, request: Request
) -> Any:
    base = store.get_plan(plan_id, payload.base_plan_version)
    if base is None:
        return _error(request, 404, "PLAN_NOT_FOUND", "找不到基準規劃版本。")
    dataset_record = store.get_dataset(base.dataset_id)
    if dataset_record is None:
        return _error(request, 404, "DATASET_NOT_FOUND", "找不到規劃資料集。")
    return _cross_vehicle_route_order_preview_payload(base, dataset_record, payload) | {
        "request_id": _request_id(request),
    }


@app.post("/api/v1/plans/{plan_id}/route-order/cross-vehicle/confirm")
def cross_vehicle_route_order_confirm(
    plan_id: str, payload: CrossVehicleRouteOrderRequest, request: Request
) -> Any:
    base = store.get_plan(plan_id, payload.base_plan_version)
    if base is None:
        return _error(request, 404, "PLAN_NOT_FOUND", "找不到基準規劃版本。")
    dataset_record = store.get_dataset(base.dataset_id)
    if dataset_record is None:
        return _error(request, 404, "DATASET_NOT_FOUND", "找不到規劃資料集。")
    candidate, reason, validation, _source, _target = _cross_vehicle_route_order_candidate(
        base, dataset_record, payload
    )
    if candidate is None or reason is not None or not validation.valid:
        return _error(
            request,
            409,
            "CROSS_VEHICLE_ROUTE_NOT_FEASIBLE",
            reason or "跨車站序未通過獨立驗證；原方案未變更。",
            validation=validation.model_dump(mode="json"),
            requires_manual_review=True,
        )
    next_version = max(store.plans.get(plan_id, {0: None})) + 1
    record = PlanRecord(
        plan_id=plan_id,
        dataset_id=base.dataset_id,
        version=next_version,
        state=base.state,
        plan=candidate,
        validation=validation,
        matrix=base.matrix,
        created_at=datetime.now(UTC).isoformat(),
    )
    store.add_plan(record)
    repository.save_plan(record)
    return _plan_payload(record) | {"request_id": _request_id(request)}


@app.get("/api/v1/plans/{plan_id}/map-data")
def get_map_data(
    plan_id: str,
    request: Request,
    version: int | None = None,
    timeline_minutes: int | None = None,
) -> Any:
    record = store.get_plan(plan_id, version)
    if record is None:
        return _error(request, 404, "PLAN_NOT_FOUND", "找不到規劃版本。")
    dataset_record = store.get_dataset(record.dataset_id)
    assert dataset_record is not None
    routes: list[dict[str, Any]] = []
    map_stage = stage_for_state(record.state)
    effective_timeline = clamp_timeline_minutes(timeline_minutes)
    progress_by_vehicle = {
        item["vehicle_id"]: item
        for item in route_progress(record.plan, timeline_minutes)
    } if map_stage == "DISPATCHED" else {}
    deviations = (
        compute_dispatch_deviations(record.plan, dataset_record.dataset, effective_timeline)
        if map_stage == "DISPATCHED"
        else None
    )
    if deviations and deviations["has_deviations"]:
        repository.append_audit(
            event_id=f"AUD-{uuid4().hex[:12].upper()}",
            event_type="DISPATCH_DEVIATION_RECORDED",
            created_at=datetime.now(UTC).isoformat(),
            plan_id=record.plan_id,
            version=record.version,
            metadata={
                "source": deviations["source"],
                "timeline_minutes": deviations["timeline_minutes"],
                "vehicle_deviation_count": len(deviations["vehicle_deviations"]),
                "zone_deviation_count": len(deviations["zone_deviations"]),
            },
        )
    google_provider = GoogleRoutesProvider(settings.google_routes_server_api_key)
    for index, route in enumerate(record.plan.routes):
        route_progress_data = progress_by_vehicle.get(route.vehicle_id, {})
        stop_status_by_order = {
            item["order_id"]: item["status"]
            for item in route_progress_data.get("stop_statuses", [])
        }
        stops = [
            stop.model_dump(include={"sequence", "order_id", "latitude", "longitude", "eta"})
            | {"status": stop_status_by_order.get(stop.order_id, "UPCOMING")}
            for stop in route.stops
        ]
        coordinates = [
            (SimulatedRouteProvider.depot_latitude, SimulatedRouteProvider.depot_longitude),
            *[(float(stop["latitude"]), float(stop["longitude"])) for stop in stops],
            (SimulatedRouteProvider.depot_latitude, SimulatedRouteProvider.depot_longitude),
        ]
        if record.matrix.provider_mode == "GOOGLE":
            try:
                encoded_polyline = google_provider.build_route_geometry(
                    coordinates, allow_fallback=False
                )
            except GoogleRoutesProviderError as exc:
                return _error(
                    request,
                    502,
                    "PROVIDER_UNAVAILABLE",
                    "Google Routes 道路幾何無法取得。",
                    provider="GOOGLE",
                    operation="computeRoutes",
                    provider_error=exc.code,
                    provider_error_category=exc.category,
                    fallback_used=False,
                    retryable=exc.code in {"GOOGLE_TIMEOUT", "GOOGLE_REQUEST_FAILED"},
                )
            is_simplified = False
        else:
            encoded_polyline = "simulated:" + ";".join(
                f"{latitude},{longitude}" for latitude, longitude in coordinates
            )
            is_simplified = True
        routes.append(
            {
                "vehicle_id": route.vehicle_id,
                "color": ["#2563EB", "#16A34A", "#EA580C", "#9333EA"][index % 4],
                "encoded_polyline": encoded_polyline,
                "is_simplified": is_simplified,
                "stops": stops,
                "completed_stops": route_progress_data.get("completed_stops", []),
                "current_stop_id": route_progress_data.get("current_stop_id"),
                "current_position": route_progress_data.get("current_position", "DEPOT-001"),
                "legs": [
                    {
                        "from_sequence": stop.sequence - 1,
                        "to_sequence": stop.sequence,
                        "distance_m": stop.leg_distance_m,
                        "duration_s": stop.leg_duration_s,
                    }
                    for stop in route.stops
                ],
            }
        )
    tdx_provider = TDXProvider(
        settings.tdx_client_id,
        settings.tdx_client_secret,
        api_base_url=settings.tdx_api_base_url,
        traffic_endpoint=settings.tdx_traffic_endpoint,
        timeout_seconds=settings.tdx_timeout_seconds,
    )
    traffic = tdx_provider.fetch_traffic()
    warnings: list[dict[str, str]] = []
    if record.matrix.provider_mode == "SIMULATED":
        warnings.append({"code": "SIMULATED_ROUTE_DATA", "message": "非 Google 即時道路資料。"})
    if traffic.warning:
        warnings.append(
            {
                "code": traffic.warning,
                "message": "TDX 路況資料目前不可用, 未以模擬事件替代。",
            }
        )
    return {
        "plan_id": plan_id,
        "version": record.version,
        "provider_mode": record.matrix.provider_mode,
        "matrix_hash": matrix_hash(record.matrix),
        "matrix_version": record.matrix.matrix_version,
        "stage": stage_for_state(record.state),
        "timeline_minutes": (
            clamp_timeline_minutes(timeline_minutes)
            if stage_for_state(record.state) == "DISPATCHED"
            else None
        ),
        "depot": {
            "depot_id": "DEPOT-001",
            "latitude": SimulatedRouteProvider.depot_latitude,
            "longitude": SimulatedRouteProvider.depot_longitude,
        },
        "routes": routes,
        "deviations": deviations,
        "traffic": {
            "mode": traffic.mode,
            "data_status": traffic.data_status,
            "events": [event.model_dump() for event in traffic.events],
            "route_risks": correlate_events_to_plan(
                dataset_record.dataset, record.plan, traffic.events
            ),
        },
        "warnings": warnings,
        "request_id": _request_id(request),
    }


def _urgent_inserted_orders(
    plan: PlanResult, order_ids: list[str], dataset: Dataset
) -> list[dict[str, Any]]:
    inserted_orders: list[dict[str, Any]] = []
    orders = {order.order_id: order for order in dataset.orders}
    vehicles = {vehicle.vehicle_id: vehicle for vehicle in dataset.vehicles}
    for order_id in order_ids:
        assignment = next(
            (
                (route.vehicle_id, stop.sequence, stop.eta)
                for route in plan.routes
                for stop in route.stops
                if stop.order_id == order_id
            ),
            None,
        )
        inserted_orders.append(
            {
                "order_id": order_id,
                "vehicle_id": assignment[0] if assignment else None,
                "sequence": assignment[1] if assignment else None,
                "eta": assignment[2] if assignment else None,
                "status": "ASSIGNED" if assignment else "UNASSIGNED",
                "responsibility": (
                    "責任區"
                    if assignment
                    and PREFERRED_VEHICLE_BY_ZONE.get(
                        orders[order_id].zone_code
                    )
                    == assignment[0]
                    else "跨區支援"
                    if assignment and assignment[0] in vehicles
                    else None
                ),
            }
        )
    return inserted_orders


def _affected_vehicle_count(diff: dict[str, Any]) -> int:
    affected_vehicles = {
        change["vehicle_id"]
        for change in diff["vehicle_load_changes"]
        if change["delta_load_kg"] != 0
    }
    affected_vehicles.update(
        change["from_vehicle_id"]
        for change in diff["sequence_changes"]
        if change["from_vehicle_id"] is not None
    )
    affected_vehicles.update(
        change["to_vehicle_id"]
        for change in diff["sequence_changes"]
        if change["to_vehicle_id"] is not None
    )
    return len(affected_vehicles)


def _urgent_option_payload(
    option: UrgentOption,
    record: PlanRecord,
    base_record: PlanRecord,
    incoming_ids: list[str],
    dataset: Dataset,
    label: str,
) -> dict[str, Any]:
    diff = option.diff
    reordered_order_count = count_reordered_orders(
        base_record.plan, option.plan, set(incoming_ids)
    )
    minimum_capacity_slack = min(
        (
            route.max_load_kg - route.planned_load_kg
            for route in option.plan.routes
            if route.order_ids
        ),
        default=0.0,
    )
    distance_delta_m = round(float(diff["total_distance_delta_m"]), 3)
    duration_delta_s = round(float(diff["total_duration_delta_s"]), 3)
    return {
        "option_id": option.option_id,
        "label": label,
        "title": option.title,
        "rationale": option.rationale,
        "mode": option.mode,
        "plan_id": record.plan_id,
        "base_version": base_record.version,
        "preview_version": record.version,
        "feasible": True,
        "selectable": True,
        "requires_human_confirmation": True,
        "inserted_orders": _urgent_inserted_orders(
            option.plan, incoming_ids, dataset
        ),
        "cost": {
            "distance_delta_m": distance_delta_m,
            "distance_delta_km": round(distance_delta_m / 1000, 1),
            "duration_delta_s": duration_delta_s,
            "duration_delta_min": round(duration_delta_s / 60, 1),
            "vehicle_change_count": len(diff["reassigned_orders"]),
            "minimum_capacity_slack_kg": round(minimum_capacity_slack, 3),
        },
        "affected_vehicle_count": _affected_vehicle_count(diff),
        "moved_order_count": len(diff["reassigned_orders"]),
        "reordered_order_count": reordered_order_count,
        "insertion": {
            "vehicle_id": option.vehicle_id,
            "sequence": option.insertion_sequence,
        },
        "after": _plan_payload(record)["summary"],
        "validator": option.validation.model_dump(mode="json"),
        "diff": diff,
    }


def _urgent_insert_preview_response(
    plan_id: str,
    base_plan_version: int,
    bundles: list[UrgentOrderBundleRequest],
    request: Request,
    *,
    include_unassignable_option: bool = False,
) -> Any:
    """Create immutable previews and deterministic options for one urgent batch."""

    base_record = store.get_plan(plan_id, base_plan_version)
    if base_record is None:
        return _error(request, 404, "PLAN_NOT_FOUND", "找不到基準規劃版本。")
    if base_record.state == "DISPATCHED":
        return _error(request, 409, "PLAN_ALREADY_DISPATCHED", "已出發的規劃不可插單。")
    dataset_record = store.get_dataset(base_record.dataset_id)
    if dataset_record is None:
        return _error(request, 404, "DATASET_NOT_FOUND", "找不到基準資料集。")
    incoming_ids = [bundle.order.order_id for bundle in bundles]
    duplicate_ids = sorted(
        {order_id for order_id in incoming_ids if incoming_ids.count(order_id) > 1}
        | {order.order_id for order in dataset_record.dataset.orders}.intersection(incoming_ids)
    )
    if duplicate_ids:
        return _error(
            request,
            422,
            "URGENT_ORDER_DUPLICATE",
            "臨時訂單編號重複，原方案未變更。",
            duplicate_order_ids=duplicate_ids,
        )
    new_orders: list[Order] = []
    for bundle in bundles:
        if bundle.order.declared_package_count != len(bundle.packages):
            return _error(
                request,
                422,
                "URGENT_ORDER_INVALID",
                f"{bundle.order.order_id} 的件數與包裹資料不一致。",
                order_id=bundle.order.order_id,
            )
        if any(package.order_id != bundle.order.order_id for package in bundle.packages):
            return _error(
                request,
                422,
                "URGENT_ORDER_INVALID",
                f"{bundle.order.order_id} 的包裹編號指向錯誤訂單。",
                order_id=bundle.order.order_id,
            )
        new_orders.append(
            Order.model_validate(
                bundle.order.model_dump() | {"packages": tuple(bundle.packages)}
            )
        )
    new_dataset = dataset_record.dataset.model_copy(
        update={
            "orders": (*dataset_record.dataset.orders, *new_orders),
            "packages": (
                *dataset_record.dataset.packages,
                *(package for bundle in bundles for package in bundle.packages),
            ),
        }
    )
    validation_report = validate_dataset(new_dataset)
    if not validation_report.is_valid:
        return _error(
            request,
            422,
            "URGENT_ORDER_INVALID",
            "插單資料未通過驗證。",
            field_errors=[error.model_dump() for error in validation_report.errors],
        )
    preview_dataset_id = f"DS-{uuid4().hex[:12].upper()}"
    try:
        preview_matrix = (
            GoogleRoutesProvider(settings.google_routes_server_api_key).extend_matrix(
                base_record.matrix,
                base_record.matrix.node_ids,
                _dataset_matrix_coordinates(dataset_record.dataset),
                (
                    "DEPOT-001",
                    *(
                        order.order_id
                        for order in sorted(new_dataset.orders, key=lambda item: item.order_id)
                    ),
                ),
                _dataset_matrix_coordinates(new_dataset),
                allow_fallback=False,
            )
            if base_record.matrix.provider_mode == "GOOGLE"
            else SimulatedRouteProvider().build(new_dataset)
        )
    except GoogleRoutesProviderError as exc:
        return _error(
            request,
            502,
            "PROVIDER_UNAVAILABLE",
            "Google Routes 即時矩陣無法取得, 無法產生插單預覽。",
            provider="GOOGLE",
            operation="computeRouteMatrix",
            provider_error=exc.code,
            provider_error_category=exc.category,
            fallback_used=False,
            retryable=exc.code in {"GOOGLE_TIMEOUT", "GOOGLE_REQUEST_FAILED"},
        )
    try:
        options = build_urgent_options(
            base_record.plan,
            new_dataset,
            preview_matrix,
            settings.solver_time_limit_seconds,
            incoming_ids,
            stage=stage_for_state(base_record.state),
            active_rules=list_dispatch_rules(include_inactive=False),
        )
        partial_plan = build_partial_urgent_plan(
            base_record.plan,
            new_dataset,
            preview_matrix,
            incoming_ids,
        )
    except Exception:
        # A provider matrix or deterministic insertion edge case must never
        # become an opaque HTTP 500. No preview record has been persisted yet.
        return _error(
            request,
            409,
            "URGENT_INSERT_UNASSIGNABLE",
            "插單無法在目前方案中合法安排，原方案未變更。",
            plan_id=plan_id,
            order_ids=incoming_ids,
            reason="PLANNER_NO_FEASIBLE_CANDIDATE",
        )
    partial_validation = validate_plan(new_dataset, partial_plan, preview_matrix)
    partial_assigned_ids = {
        order_id
        for route in partial_plan.routes
        for order_id in route.order_ids
    }
    unassigned_incoming = [
        order_id for order_id in incoming_ids if order_id not in partial_assigned_ids
    ]
    if unassigned_incoming or not options:
        route_by_vehicle = {route.vehicle_id: route for route in partial_plan.routes}
        vehicle_by_id = {
            vehicle.vehicle_id: vehicle for vehicle in dataset_record.dataset.vehicles
        }
        order_by_id = {order.order_id: order for order in new_orders}
        readable_reasons: dict[str, str] = {}
        for order_id in unassigned_incoming:
            order = order_by_id[order_id]
            eligible = [
                vehicle
                for vehicle in vehicle_by_id.values()
                if vehicle.status == VehicleStatus.AVAILABLE
                and order.zone_code in vehicle.service_zone_codes
            ]
            if not eligible:
                readable_reasons[order_id] = "SERVICE_ZONE_UNAVAILABLE"
            elif all(
                order.total_weight_kg
                > vehicle.max_load_kg
                - route_by_vehicle[vehicle.vehicle_id].planned_load_kg
                for vehicle in eligible
            ):
                readable_reasons[order_id] = "CAPACITY_LIMIT"
            else:
                readable_reasons[order_id] = "TIME_OR_ROUTE_CONFLICT"
        if include_unassignable_option:
            diff = compute_plan_diff(base_record.plan, partial_plan)
            inserted_orders = _urgent_inserted_orders(
                partial_plan, incoming_ids, new_dataset
            )
            return {
                "plan_id": plan_id,
                "base_version": base_record.version,
                "preview_version": base_record.version,
                "feasible": False,
                "requires_human_confirmation": True,
                "mode": "UNASSIGNABLE",
                "rejection_reason": "NO_LEGAL_BATCH_ROUTE_INSERTION",
                "affected_vehicle_count": _affected_vehicle_count(diff),
                "moved_order_count": len(diff["reassigned_orders"]),
                "before": _plan_payload(base_record)["summary"],
                "after": _plan_payload(base_record)["summary"],
                "comparison": {
                    "base_algorithm": base_record.plan.algorithm,
                    "preview_algorithm": partial_plan.algorithm,
                    "base_dataset_hash": dataset_hash(dataset_record.dataset),
                    "preview_dataset_hash": dataset_hash(new_dataset),
                },
                "validator": partial_validation.model_dump(mode="json"),
                "provider_mode": preview_matrix.provider_mode,
                "matrix_hash": matrix_hash(preview_matrix),
                "matrix_reused": base_record.matrix.provider_mode == preview_matrix.provider_mode,
                "inserted_orders": inserted_orders,
                "options": [
                    {
                        "option_id": "UNASSIGNABLE",
                        "label": "方案 A",
                        "title": "排不進去",
                        "rationale": "目前沒有同時符合載重、服務區域與時段的合法安排。",
                        "mode": "UNASSIGNABLE",
                        "plan_id": plan_id,
                        "base_version": base_record.version,
                        "preview_version": base_record.version,
                        "feasible": False,
                        "selectable": False,
                        "requires_human_confirmation": True,
                        "inserted_orders": inserted_orders,
                        "unassigned_orders": unassigned_incoming,
                        "unassigned_reasons": readable_reasons,
                        "cost": {
                            "distance_delta_m": None,
                            "distance_delta_km": None,
                            "duration_delta_s": None,
                            "duration_delta_min": None,
                            "vehicle_change_count": 0,
                            "minimum_capacity_slack_kg": None,
                        },
                        "affected_vehicle_count": _affected_vehicle_count(diff),
                        "moved_order_count": len(diff["reassigned_orders"]),
                        "after": _plan_payload(base_record)["summary"],
                        "validator": partial_validation.model_dump(mode="json"),
                        "diff": diff,
                    }
                ],
                "diff": {
                    "inserted_order_ids": incoming_ids,
                    "inserted_order_id": incoming_ids[0] if len(incoming_ids) == 1 else None,
                    **diff,
                },
                "warnings": [],
                "request_id": _request_id(request),
            }
        return _error(
            request,
            409,
            "URGENT_INSERT_UNASSIGNABLE",
            "一張或多張臨時訂單目前無法合法安排，原方案未變更。",
            plan_id=plan_id,
            order_ids=incoming_ids,
            unassigned_order_ids=unassigned_incoming,
            unassigned_reasons=readable_reasons,
        )
    if not options:
        return _error(
            request,
            409,
            "URGENT_INSERT_UNASSIGNABLE",
            "插單無法在目前方案中合法安排，原方案未變更。",
            plan_id=plan_id,
            order_ids=incoming_ids,
            reason="NO_VALID_OPTION",
        )
    preview_version = max(store.plans.get(plan_id, {0: None})) + 1
    preview_records: list[PlanRecord] = []
    created_at = datetime.now(UTC).isoformat()
    preview_state = (
        base_record.state if base_record.state in {"LOADED", "DISPATCHED"} else "PROPOSED"
    )
    for offset, option in enumerate(options):
        preview_plan_for_state = option.plan.model_copy(update={"state": preview_state})
        preview_record = PlanRecord(
            plan_id=plan_id,
            dataset_id=preview_dataset_id,
            version=preview_version + offset,
            state=preview_state,
            plan=preview_plan_for_state,
            validation=option.validation,
            matrix=preview_matrix,
            created_at=created_at,
        )
        preview_records.append(preview_record)
    primary_record = preview_records[0]
    preview_dataset_record = DatasetRecord(
        dataset_id=preview_dataset_id,
        dataset=new_dataset,
        validation=validation_report,
        matrix=preview_matrix,
        created_at=created_at,
    )
    store.add_dataset(preview_dataset_record)
    repository.save_dataset(
        preview_dataset_id,
        new_dataset,
        validation_report,
        preview_matrix,
        created_at,
    )
    for preview_record in preview_records:
        store.add_plan(preview_record, make_current=False)
        repository.save_plan(preview_record, make_current=False)
    preview_plan = primary_record.plan
    preview_validation = primary_record.validation
    diff = compute_plan_diff(base_record.plan, preview_plan)
    preview_warnings: list[dict[str, str]] = []
    if preview_matrix.provider_mode == "SIMULATED":
        preview_warnings.append(
            {
                "code": preview_matrix.warning or "SIMULATED_ROUTE_DATA",
                "message": "目前使用可重現的模擬距離與路線資料, 非 Google 即時資料。",
            }
        )
    elif preview_matrix.warning:
        preview_warnings.append(
            {"code": preview_matrix.warning, "message": "路線 provider 回傳警告。"}
        )
    inserted_orders = _urgent_inserted_orders(preview_plan, incoming_ids, new_dataset)
    option_payloads = [
        _urgent_option_payload(
            option,
            preview_records[index],
            base_record,
            incoming_ids,
            new_dataset,
            f"方案 {chr(ord('A') + index)}",
        )
        for index, option in enumerate(options)
    ]
    return {
        "plan_id": plan_id,
        "base_version": base_record.version,
        "preview_version": primary_record.version,
        "feasible": True,
        "requires_human_confirmation": True,
        "mode": "INSERTION",
        "affected_vehicle_count": _affected_vehicle_count(diff),
        "moved_order_count": len(diff["reassigned_orders"]),
        "before": _plan_payload(base_record)["summary"],
        "after": _plan_payload(primary_record)["summary"],
        "comparison": {
            "base_algorithm": base_record.plan.algorithm,
            "preview_algorithm": preview_plan.algorithm,
            "base_dataset_hash": dataset_hash(dataset_record.dataset),
            "preview_dataset_hash": dataset_hash(new_dataset),
        },
        "validator": preview_validation.model_dump(mode="json"),
        "provider_mode": preview_matrix.provider_mode,
        "matrix_hash": matrix_hash(preview_matrix),
        "matrix_reused": base_record.matrix.provider_mode == preview_matrix.provider_mode,
        "matrix_elements_added": (
            len(new_orders) * len(preview_matrix.node_ids)
            + len(base_record.matrix.node_ids) * len(new_orders)
            if preview_matrix.provider_mode == "GOOGLE"
            else 0
        ),
        "inserted_orders": inserted_orders,
        "options": option_payloads,
        "diff": {
            "inserted_order_ids": incoming_ids,
            "inserted_order_id": incoming_ids[0] if len(incoming_ids) == 1 else None,
            **diff,
        },
        "warnings": preview_warnings,
        "request_id": _request_id(request),
    }


@app.post("/api/v1/plans/{plan_id}/urgent-insert/preview")
def urgent_insert_preview(plan_id: str, payload: UrgentInsertRequest, request: Request) -> Any:
    bundles = [UrgentOrderBundleRequest(order=payload.order, packages=payload.packages)]
    response = _urgent_insert_preview_response(
        plan_id,
        payload.base_plan_version,
        bundles,
        request,
    )
    _sync_direct_urgent_preview_session(
        request, plan_id, payload.base_plan_version, bundles, response
    )
    return response


@app.post("/api/v1/plans/{plan_id}/urgent-insert/batch-preview")
def urgent_batch_insert_preview(
    plan_id: str,
    payload: UrgentBatchInsertRequest,
    request: Request,
    include_unassignable_option: bool = False,
) -> Any:
    response = _urgent_insert_preview_response(
        plan_id,
        payload.base_plan_version,
        payload.orders,
        request,
        include_unassignable_option=include_unassignable_option,
    )
    _sync_direct_urgent_preview_session(
        request, plan_id, payload.base_plan_version, payload.orders, response
    )
    return response


@app.post("/api/v1/plans/{plan_id}/confirm")
def confirm_plan(plan_id: str, payload: ConfirmRequest, request: Request) -> Any:
    record = store.get_plan(plan_id, payload.version)
    if record is None:
        return _error(request, 404, "PLAN_NOT_FOUND", "找不到規劃版本。")
    dispatched_scope_preview = False
    if record.state == "DISPATCHED":
        current = store.get_plan(plan_id)
        session_id = request.headers.get("X-Dispatch-Session")
        session = agent_sessions.get(session_id) if session_id else None
        if session is None and session_id:
            persisted_session = repository.load_agent_session(session_id)
            session = _session_from_payload(persisted_session) if persisted_session else None
        frozen_ids = session.frozen_stop_ids if session is not None else ()
        if (
            current is not None
            and current.state == "DISPATCHED"
            and record.version > current.version
            and record.dataset_id == current.dataset_id
            and record.plan.algorithm == "ORTOOLS"
            and record.validation.valid
        ):
            current_routes = {route.vehicle_id: route for route in current.plan.routes}
            candidate_routes = {route.vehicle_id: route for route in record.plan.routes}
            dispatched_scope_preview = set(current_routes) == set(candidate_routes)
            if dispatched_scope_preview:
                for vehicle_id, current_route in current_routes.items():
                    candidate_route = candidate_routes[vehicle_id]
                    if set(candidate_route.order_ids) != set(current_route.order_ids):
                        dispatched_scope_preview = False
                        break
                    frozen_prefix = []
                    for order_id in current_route.order_ids:
                        if order_id in frozen_ids:
                            frozen_prefix.append(order_id)
                        else:
                            break
                    if candidate_route.order_ids[: len(frozen_prefix)] != frozen_prefix:
                        dispatched_scope_preview = False
                        break
            if not dispatched_scope_preview:
                return _error(
                    request,
                    409,
                    "PLAN_ALREADY_DISPATCHED",
                    "已出發的規劃只能確認保留已送站點與車輛指派的途中調整。",
                )
        else:
            return _error(request, 409, "PLAN_ALREADY_DISPATCHED", "已出發的規劃不可再次確認。")
    known_unassigned = set(record.plan.unassigned_orders)
    preserves_existing_unassigned = False
    if known_unassigned:
        previous_versions = store.plans.get(plan_id, {})
        preserves_existing_unassigned = any(
            version < record.version
            and known_unassigned.issubset(set(previous.plan.unassigned_orders))
            for version, previous in previous_versions.items()
        )
    if not dispatched_scope_preview and (
        record.state != "PROPOSED"
        or record.plan.algorithm != "ORTOOLS"
        or not record.validation.valid
        or (not record.plan.complete and not preserves_existing_unassigned)
    ):
        return _error(request, 409, "PLAN_NOT_CONFIRMABLE", "規劃尚未通過驗證或狀態不允許確認。")
    if not dispatched_scope_preview:
        record.state = "CONFIRMED"
    # A confirmed version becomes the current read/continuation pointer. Preview
    # versions remain immutable and never become current before this checkpoint.
    store.current_versions[plan_id] = record.version
    repository.set_current_version(plan_id, record.version)
    repository.update_plan_state(plan_id, record.version, record.state)
    repository.append_audit(
        f"AUD-{uuid4().hex[:12].upper()}",
        "PLAN_EN_ROUTE_ADJUSTED" if dispatched_scope_preview else "PLAN_CONFIRMED",
        datetime.now(UTC).isoformat(),
        plan_id,
        record.version,
    )
    session_id = request.headers.get("X-Dispatch-Session")
    if session_id:
        session = agent_sessions.get(session_id)
        if session is None:
            persisted_session = repository.load_agent_session(session_id)
            session = _session_from_payload(persisted_session) if persisted_session else None
            if session is not None:
                agent_sessions[session_id] = session
        if session is not None and session.plan_id == plan_id:
            session.urgent_workflow = UrgentWorkflowState().model_dump(mode="json")
            session.last_preview_version = None
            session.plan_version = record.version
            _save_agent_session(session_id, session)
    return _plan_payload(record) | {
        "audit_event_id": f"AUD-{uuid4().hex[:12].upper()}",
        "request_id": _request_id(request),
    }


@app.post("/api/v1/plans/{plan_id}/dispatch")
def dispatch_plan(plan_id: str, payload: DispatchRequest, request: Request) -> Any:
    if not settings.dispatch_enabled:
        return _error(
            request,
            403,
            "DISPATCH_DISABLED",
            "本測試環境已停用 Dispatch；請由調度員在外部流程另行處理。",
        )
    record = store.get_plan(plan_id, payload.version)
    if record is None:
        return _error(request, 404, "PLAN_NOT_FOUND", "找不到規劃版本。")
    if record.state != "CONFIRMED":
        return _error(request, 409, "PLAN_NOT_CONFIRMABLE", "只有已確認版本可以標記出發。")
    record.state = "DISPATCHED"
    repository.update_plan_state(plan_id, record.version, record.state)
    repository.append_audit(
        f"AUD-{uuid4().hex[:12].upper()}",
        "PLAN_DISPATCHED",
        datetime.now(UTC).isoformat(),
        plan_id,
        record.version,
    )
    return _plan_payload(record) | {
        "audit_event_id": f"AUD-{uuid4().hex[:12].upper()}",
        "request_id": _request_id(request),
    }


@app.post("/api/v1/plans/{plan_id}/load")
def load_plan(plan_id: str, payload: LoadingRequest, request: Request) -> Any:
    """Start the manual loading stage without calling formal dispatch."""
    record = store.get_plan(plan_id, payload.version)
    if record is None:
        return _error(request, 404, "PLAN_NOT_FOUND", "找不到要裝車的規劃版本。")
    if record.state == "LOADED":
        return _error(request, 409, "PLAN_ALREADY_LOADED", "這個版本已經在上車後階段。")
    if record.state == "DISPATCHED":
        return _error(request, 409, "PLAN_ALREADY_DISPATCHED", "已發車的規劃不能回到裝車階段。")
    if (
        record.state not in {"PROPOSED", "CONFIRMED"}
        or not record.validation.valid
    ):
        return _error(request, 409, "PLAN_NOT_LOADABLE", "方案尚未通過完整性驗證，不能開始裝車。")
    record.state = "LOADED"
    repository.update_plan_state(plan_id, record.version, record.state)
    repository.append_audit(
        f"AUD-{uuid4().hex[:12].upper()}",
        "PLAN_LOADING_STARTED",
        datetime.now(UTC).isoformat(),
        plan_id,
        record.version,
        {"dispatcher_reference": payload.dispatcher_reference},
    )
    return _plan_payload(record) | {"request_id": _request_id(request)}


@app.post("/api/v1/plans/{plan_id}/simulate-departure")
def simulate_departure(
    plan_id: str, payload: SimulatedDepartureRequest, request: Request
) -> Any:
    """Advance the local demo state to F5; this is never formal dispatch."""
    record = store.get_plan(plan_id, payload.version)
    if record is None:
        return _error(request, 404, "PLAN_NOT_FOUND", "找不到要模擬出發的規劃版本。")
    if record.state == "DISPATCHED":
        return _error(request, 409, "PLAN_ALREADY_DISPATCHED", "這個版本已經在已發車階段。")
    if record.state != "LOADED":
        return _error(request, 409, "PLAN_NOT_LOADED", "請先完成開始裝車，才能進入已發車階段。")
    record.state = "DISPATCHED"
    repository.update_plan_state(plan_id, record.version, record.state)
    repository.append_audit(
        f"AUD-{uuid4().hex[:12].upper()}",
        "PLAN_SIMULATED_DEPARTURE",
        datetime.now(UTC).isoformat(),
        plan_id,
        record.version,
        {"formal_dispatch_called": False},
    )
    return _plan_payload(record) | {
        "notice": "已進入本地示範的已發車階段；正式 Dispatch 仍維持停用。",
        "request_id": _request_id(request),
    }


@app.get("/api/v1/providers/status")
def provider_status(request: Request) -> dict[str, Any]:
    status_map = settings.credential_status()
    tdx_status = TDXProvider(
        settings.tdx_client_id,
        settings.tdx_client_secret,
        api_base_url=settings.tdx_api_base_url,
        traffic_endpoint=settings.tdx_traffic_endpoint,
        timeout_seconds=settings.tdx_timeout_seconds,
    ).status()
    return {
        "providers": [
            {"name": "simulated_routes", "enabled": True, "status": "healthy", "mode": "SIMULATED"},
            {
                "name": "google_routes",
                "enabled": settings.google_routes_enabled
                and status_map["GOOGLE_ROUTES_SERVER_API_KEY"] == "CONFIGURED",
                "status": provider_runtime_state["google_routes"]
                if settings.google_routes_enabled
                and status_map["GOOGLE_ROUTES_SERVER_API_KEY"] == "CONFIGURED"
                else "disabled",
                "mode": "GOOGLE"
                if settings.google_routes_enabled
                and status_map["GOOGLE_ROUTES_SERVER_API_KEY"] == "CONFIGURED"
                else "UNAVAILABLE",
            },
            {"name": "tdx", **tdx_status.model_dump()},
            {
                "name": "openai",
                "enabled": status_map["OPENAI_API_KEY"] == "CONFIGURED",
                "status": provider_runtime_state["openai"]
                if status_map["OPENAI_API_KEY"] == "CONFIGURED"
                else "disabled",
                "mode": "OPENAI" if status_map["OPENAI_API_KEY"] == "CONFIGURED" else "UNAVAILABLE",
            },
        ],
        "request_id": _request_id(request),
    }


def _urgent_validation_errors(response: JSONResponse) -> list[dict[str, Any]]:
    try:
        raw_body = response.body
        if isinstance(raw_body, memoryview):
            raw_body = raw_body.tobytes()
        body = json.loads(raw_body)
    except (TypeError, json.JSONDecodeError):
        return []
    error = body.get("error") if isinstance(body, dict) else None
    field_errors = error.get("field_errors") if isinstance(error, dict) else None
    if not isinstance(field_errors, list):
        return []
    return [item for item in field_errors if isinstance(item, dict)]


def _urgent_validation_order_id(
    response: JSONResponse, orders: list[dict[str, Any]]
) -> str | None:
    errors = _urgent_validation_errors(response)
    for error in errors:
        path = error.get("path")
        if not isinstance(path, str):
            continue
        for order in orders:
            order_id = order.get("order_id")
            if isinstance(order_id, str) and path.startswith(f"orders.{order_id}."):
                return order_id
    return None


def _urgent_workflow_message(
    stage: str,
    orders: list[dict[str, Any]],
    missing: list[dict[str, Any]],
    duplicate_ids: list[str],
    preview: dict[str, Any] | None = None,
) -> str:
    if stage == "CANCELLED":
        return "已取消這次臨時插單，原方案沒有變更。"
    if stage == "BLOCKED" and duplicate_ids:
        return f"訂單編號重複：{'、'.join(duplicate_ids)}。請修改編號後再試，原方案沒有變更。"
    if stage == "BLOCKED":
        return "不能跳過資料檢查、插單預覽或人工確認；原方案沒有變更。"
    if missing:
        missing_names = {str(field) for item in missing for field in item["missing_fields"]}
        display_fields: list[str] = []
        if "order_id" in missing_names:
            display_fields.append("訂單編號")
        if {"location_label", "city", "district"} & missing_names:
            display_fields.append("配送地點")
        if {"latitude", "longitude"} & missing_names:
            display_fields.append("座標")
        if "zone_code" in missing_names:
            display_fields.append("配送區域")
        if "package_weight_kg" in missing_names:
            display_fields.append("重量")
        if "declared_package_count" in missing_names:
            display_fields.append("件數")
        if "time_slot" in missing_names:
            display_fields.append("配送時段")
        # One unbroken list. Wrapping it after the third item used to split a
        # single sentence mid-way, so a line could end on a field name and the
        # next line start with the full stop.
        # Count-neutral on purpose. 「三張急單」 still produces a single draft
        # at this point, so saying 「這張急單」 contradicted what the dispatcher
        # had just said, and saying 「這 1 張」 would be just as wrong.
        return (
            f"還缺少幾個欄位才能算：{'、'.join(display_fields)}。\n"
            "每一張都照這個順序補齊給我。"
        )
    if stage == "REVIEW_READY":
        # One order per line. Run together with punctuation separators this was
        # a wall of text a dispatcher could not scan, and it leaked the raw
        # MORNING / AFTERNOON enum into a Chinese sentence.
        summaries = []
        for order in orders:
            count = int(order["declared_package_count"])
            unit_weight = float(order["package_weight_kg"])
            resolved_priority = UrgentOrderDraft.resolve_priority(order.get("priority"))
            priority = "　急件" if resolved_priority == "HIGH" else ""
            summaries.append(
                f"{order['order_id']}　{order['location_label']}（{order['zone_code']}）"
                f"　{unit_weight:g} 公斤 × {count} 件"
                f"　{slot_sentence(order['time_slot'])}{priority}"
            )
            # Anything the application filled in has to be on the card. A city
            # or district the dispatcher never said is exactly what the human
            # confirmation step exists to catch, and it cannot be caught if it
            # is not shown.
            derived = {str(name) for name in (order.get("derived_fields") or [])}
            filled = "".join(
                str(order[key])
                for key in ("city", "district")
                if key in derived and order.get(key)
            )
            if filled:
                summaries.append(
                    f"　（{filled} 是我依 {order['zone_code']} 區補的，不對就跟我說）"
                )
        count_word = "這張" if len(orders) == 1 else f"這 {len(orders)} 張"
        return (
            f"{count_word}我記下來了，確認一下：\n\n"
            + "\n".join(summaries)
            + "\n\n沒問題就按【產生插單預覽】。要改哪一張直接跟我說。"
        )
    if stage == "PREVIEW_READY" and preview is not None:
        assignments = []
        for item in preview.get("inserted_orders", []):
            if item.get("status") == "ASSIGNED":
                assignments.append(
                    f"{item['order_id']} 排進{vehicle_label(item['vehicle_id'])}"
                    f"第 {item['sequence']} 站"
                )
        diff = preview.get("diff", {})
        option_count = sum(
            1
            for option in preview.get("options", [])
            if option.get("option_id") != "UNASSIGNABLE" and option.get("feasible") is True
        )
        assignment_text = "、".join(assignments) if assignments else "沒有可以安排的位置"
        moved = int(preview.get("moved_order_count", 0) or 0)
        # Metres and seconds are how the solver keeps score; a dispatcher thinks
        # in kilometres and minutes, so convert before the number is shown.
        distance_km = float(diff.get("total_distance_delta_m", 0) or 0) / 1000
        duration_min = float(diff.get("total_duration_delta_s", 0) or 0) / 60
        lines = [f"試算結果：{assignment_text}。"]
        lines.append(
            f"全隊多跑 {distance_km:+.1f} 公里、多花 {duration_min:+.1f} 分鐘；"
            + ("既有訂單都不用換車。" if moved == 0 else f"有 {moved} 張既有訂單要換車。")
        )
        if option_count:
            lines.append("")
            lines.append(
                f"下面有 {option_count} 種安排方式，選一張看細節。這只是預覽，還沒套用。"
                if option_count > 1
                else "下面只有一種可行的安排。這只是預覽，還沒套用。"
            )
        else:
            lines.append("")
            lines.append("沒有可行的安排，畫面會標出需要人工處理的訂單。")
        return "\n".join(lines)
    return "已收到臨時插單要求，原方案尚未變更。"


def _nearest_order_in_zone(
    dataset: Dataset, zone_code: str, latitude: float, longitude: float
) -> Order | None:
    """Find the workbook order closest to a coordinate inside one zone.

    Used to fill a missing city or district. Taking ``covered_districts[0]``
    instead put an order in 三重 merely because that name sorts first among the
    four districts Z4 covers, and the guess never appeared on the confirmation
    card, so no human could catch it. An existing order at almost the same spot
    is a value the workbook actually records for that location.
    """
    candidates = [order for order in dataset.orders if order.zone_code == zone_code]
    if not candidates:
        return None
    return min(
        candidates,
        key=lambda order: (order.latitude - latitude) ** 2
        + (order.longitude - longitude) ** 2,
    )


def _normalize_urgent_order_context(
    understanding: UrgentUnderstanding, dataset: Dataset
) -> UrgentUnderstanding:
    """Resolve a uniquely covered district to the dataset's canonical zone.

    The Agent still extracts the user's words into typed fields. This helper
    only applies deterministic dataset facts, so a district such as 信義區 can
    be validated against the workbook's canonical 信義 value without asking
    the dispatcher to repeat a zone code that the data already determines.
    """
    normalized_orders: list[UrgentOrderDraft] = []
    for draft in understanding.orders:
        zone = next(
            (item for item in dataset.zones if item.zone_code == draft.zone_code),
            None,
        )
        coordinate_complete = draft.latitude is not None and draft.longitude is not None
        derived_updates: dict[str, Any] = {}
        derived_fields = list(draft.derived_fields)
        if coordinate_complete and zone is not None:
            assert draft.latitude is not None and draft.longitude is not None
            neighbour = _nearest_order_in_zone(
                dataset, zone.zone_code, draft.latitude, draft.longitude
            )
            if not draft.city:
                city = neighbour.city if neighbour else None
                if not city and len(zone.covered_cities) == 1:
                    city = zone.covered_cities[0]
                if city:
                    derived_updates["city"] = city
                    derived_fields.append("city")
            if not draft.district:
                district = neighbour.district if neighbour else None
                if not district and len(zone.covered_districts) == 1:
                    district = zone.covered_districts[0]
                if district:
                    derived_updates["district"] = district
                    derived_fields.append("district")
            if not draft.location_label:
                derived_updates["location_label"] = f"{zone.zone_name}配送點"
                derived_fields.append("location_label")
            if derived_fields:
                derived_updates["derived_fields"] = list(dict.fromkeys(derived_fields))
        if derived_updates:
            draft = draft.model_copy(update=derived_updates)
        if not draft.district:
            normalized_orders.append(draft)
            continue
        raw_district = draft.district.strip()
        candidates: list[tuple[Zone, str]] = []
        for zone in dataset.zones:
            if draft.zone_code and zone.zone_code != draft.zone_code:
                continue
            if draft.city and draft.city not in zone.covered_cities:
                continue
            for covered_district in zone.covered_districts:
                district_matches = raw_district == covered_district or (
                    raw_district.endswith("區")
                    and raw_district.removesuffix("區") == covered_district
                )
                if district_matches:
                    candidates.append((zone, covered_district))
        unique_candidates = {
            (zone.zone_code, district): (zone, district)
            for zone, district in candidates
        }
        if len(unique_candidates) != 1:
            normalized_orders.append(draft)
            continue
        zone, canonical_district = next(iter(unique_candidates.values()))
        updates: dict[str, Any] = {"district": canonical_district}
        if draft.zone_code is None:
            updates["zone_code"] = zone.zone_code
            updates["derived_fields"] = [*draft.derived_fields, "zone_code"]
        normalized_orders.append(draft.model_copy(update=updates))
    return understanding.model_copy(update={"orders": normalized_orders})


def _json_message(value: Any) -> str | None:
    """Extract a user-facing message from a JSON-shaped model/tool string."""
    if not isinstance(value, str):
        return None
    stripped = value.strip()
    if not stripped.startswith(("{", "[")):
        return None
    try:
        parsed = json.loads(stripped)
    except (TypeError, ValueError):
        return None
    if isinstance(parsed, dict) and isinstance(parsed.get("message"), str):
        return parsed["message"].strip() or None
    if isinstance(parsed, list):
        for entry in parsed:
            if isinstance(entry, dict) and isinstance(entry.get("message"), str):
                return entry["message"].strip() or None
    return None


def _message_number(value: Any) -> str | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return f"{float(value):g}"


def _message_ids(value: Any) -> str:
    if not isinstance(value, list):
        return ""
    return "、".join(str(item) for item in value if isinstance(item, str))


def _eta_clock(value: Any) -> str | None:
    if not isinstance(value, str):
        return None
    try:
        return datetime.fromisoformat(value).strftime("%H:%M")
    except ValueError:
        return value


def _multiple_urgent_preview_sentence(item: dict[str, Any], order_ids: str) -> str:
    """Describe a whole-plan urgent preview in the words the numbers support.

    ``preview_multiple_urgent_insert`` returns a before/after plan diff. Nothing
    on screen renders from it — the option cards only come from the urgent
    workflow — so this sentence is the entire result the dispatcher gets. It
    therefore states what changed and what it cost instead of pointing at a
    card that will never appear.
    """
    before = item.get("before") if isinstance(item.get("before"), dict) else {}
    after = item.get("after") if isinstance(item.get("after"), dict) else {}
    diff = item.get("diff") if isinstance(item.get("diff"), dict) else {}
    inserted = _message_number(
        (after or {}).get("assigned_order_count", 0) - (before or {}).get("assigned_order_count", 0)
    )
    extra_km = (after or {}).get("total_distance_m", 0) - (before or {}).get("total_distance_m", 0)
    extra_min = (after or {}).get("total_duration_s", 0) - (before or {}).get("total_duration_s", 0)
    reassigned = (diff or {}).get("reassigned_orders") or []
    cost = (
        f"全隊多跑 {extra_km / 1000:.1f} 公里、多花 {extra_min / 60:.0f} 分鐘"
        if extra_km or extra_min
        else "不用多繞路，也不會多花時間"
    )
    moved = (
        f"有 {len(reassigned)} 張既有訂單要換車"
        if reassigned
        else "既有訂單都不用換車"
    )
    unassigned = (after or {}).get("unassigned_orders") or []
    tail = (
        f"\n其中 {_message_ids(unassigned)} 排不進去，要人工處理。" if unassigned else ""
    )
    return (
        f"{order_ids or '這批急單'} 都排得進去，這次多排 {inserted or '0'} 站。\n"
        f"{cost}；{moved}。{tail}\n"
        "這只是試算，原方案還沒變更。"
    )


def _operator_tool_template(tool: Any, item: dict[str, Any]) -> str:
    """Turn one deterministic evidence record into a Traditional-Chinese sentence."""
    if tool == "assistant_help":
        topic = cast(str | None, item.get("topic"))
        topics = {
            "IDENTITY": "我是配送調度助理。你把今天的訂單丟給我，我排出每台車要走的路線。",
            "CAPABILITIES": "我可以整理訂單、安排車輛、檢查路線並預覽方案；方案都由你確認。",
            "DATA_REQUIREMENTS": "建立方案前需要訂單、包裹、車輛與配送區域資料。",
            "CAPACITY_RULES": (
                "載重會依每件包裹重量加總，再和車輛上限比較；數字以方案計算結果為準。"
            ),
            "URGENT_INSERTION": "急單會先整理成摘要，再試算可行的插入位置，人工確認後才套用。",
        }
        if topic is None:
            return "我是配送調度助理，會用可核對的資料協助安排配送。"
        return topics.get(topic, "我是配送調度助理，會用可核對的資料協助安排配送。")
    if tool == "request_missing_fields":
        labels = {
            "order_id": "訂單編號",
            "zone_code": "配送區域",
            "city": "城市",
            "district": "行政區",
            "location_label": "地點名稱",
            "latitude": "緯度",
            "longitude": "經度",
            "time_slot": "配送時段",
            "declared_package_count": "包裹件數",
            "packages": "包裹重量資料",
        }
        fields = item.get("missing_fields")
        names = (
            "、".join(labels.get(str(field), str(field)) for field in fields)
            if isinstance(fields, list)
            else "必要配送欄位"
        )
        return f"目前還不能進行安全的預覽，請補充：{names}。"
    if tool == "begin_urgent_insertion":
        orders = item.get("orders")
        count = len(orders) if isinstance(orders, list) else 0
        return f"已收到 {count} 張臨時訂單資料，接下來會檢查欄位完整性。"
    if tool == "prepare_confirmation":
        return "方案仍需由調度人員在畫面上確認；助理不會直接正式派車。"
    if tool == "reject_unsupported_change":
        return "這個我不能改；目前只支援方案卡列出的六種配送調整。"
    if tool == "plan_dispatch":
        if item.get("status") == "INFEASIBLE":
            unassigned = _message_ids(item.get("unassigned_orders"))
            reason_map = item.get("unassigned_reasons")
            reason_text = ""
            if isinstance(reason_map, dict):
                details = [
                    f"{order_id}：{reason}"
                    for order_id, reason in reason_map.items()
                    if isinstance(order_id, str) and isinstance(reason, str)
                ]
                reason_text = f"；原因：{'、'.join(details)}" if details else ""
            return f"目前排不出來：{unassigned or '沒有訂單'} 無法指派{reason_text}。"
        assigned = _message_number(item.get("assigned_order_count"))
        total = _message_number(item.get("total_order_count"))
        if assigned is not None:
            total_text = f"/{total}" if total is not None else ""
            unassigned = _message_ids(item.get("unassigned_orders"))
            suffix = f"；{unassigned} 未安排" if unassigned else "；全部訂單都有安排"
            return f"配送方案試算完成：已安排 {assigned}{total_text} 張訂單，{suffix}。"
        return "配送方案試算完成，請檢查後再確認。"
    if tool == "highest_load_vehicle":
        vehicle_id = item.get("vehicle_id")
        load = _message_number(item.get("planned_load_kg"))
        limit = _message_number(item.get("max_load_kg"))
        if isinstance(vehicle_id, str) and load is not None:
            suffix = f"，載重上限 {limit} kg" if limit is not None else ""
            return f"{vehicle_label(vehicle_id)}目前計畫載重 {load} kg{suffix}。"
        return "目前沒有可查詢的車輛載重。"
    if tool == "lowest_load_vehicle":
        vehicle_id = item.get("vehicle_id")
        remaining = _message_number(item.get("remaining_capacity_kg"))
        if isinstance(vehicle_id, str) and remaining is not None:
            return f"{vehicle_label(vehicle_id)}目前剩餘容量 {remaining} kg。"
        return "目前沒有可查詢的車輛容量。"
    if tool == "vehicle_load":
        vehicle_id = item.get("vehicle_id")
        load = _message_number(item.get("planned_load_kg"))
        limit = _message_number(item.get("max_load_kg"))
        utilization = _message_number(item.get("load_utilization"))
        remaining = _message_number(item.get("remaining_capacity_kg"))
        if isinstance(vehicle_id, str) and all(
            value is not None for value in (load, limit, utilization, remaining)
        ):
            return (
                f"{vehicle_label(vehicle_id)}目前計畫載重 {load} kg，上限 {limit} kg，"
                f"使用率 {float(item['load_utilization']) * 100:g}%，剩餘容量 {remaining} kg。"
            )
        if isinstance(vehicle_id, str):
            return f"找不到車輛 {vehicle_id}，請從目前車輛清單選擇。"
        return "目前沒有可查詢的車輛載重。"
    if tool == "inspect_plan_overview":
        assigned = _message_number(item.get("assigned_order_count"))
        total = _message_number(item.get("total_order_count"))
        unassigned = _message_ids(item.get("unassigned_orders"))
        vehicle_loads = item.get("vehicles")
        load_summary: list[str] = []
        if isinstance(vehicle_loads, list):
            for vehicle in vehicle_loads:
                if not isinstance(vehicle, dict):
                    continue
                vehicle_id = vehicle.get("vehicle_id")
                load = _message_number(vehicle.get("planned_load_kg"))
                limit = _message_number(vehicle.get("max_load_kg"))
                if isinstance(vehicle_id, str) and load is not None and limit is not None:
                    # _message_number 回傳的已經是格式化過的字串, 再套 :g 會拋
                    # ValueError: Unknown format code 'g' for object of type 'str',
                    # 讓「現在的方案長什麼樣」「今天成效如何」這類問句全部回 502。
                    load_summary.append(f"{vehicle_label(vehicle_id)} {load}/{limit} kg")
        if assigned is not None and total is not None:
            suffix = f"；未安排：{unassigned}" if unassigned else "；目前沒有未安排訂單"
            if load_summary:
                suffix += "；" + "、".join(load_summary)
            return f"目前方案已安排 {assigned}/{total} 張訂單{suffix}。"
        return "目前已有配送方案，可查看各車分配、載重與未安排訂單。"
    if tool == "inspect_dispatch_deviations":
        vehicle_items = item.get("vehicle_deviations")
        zone_items = item.get("zone_deviations")
        if item.get("view") == "SUGGESTIONS":
            suggestions = item.get("suggestions")
            if isinstance(suggestions, list) and suggestions:
                suggestion_messages = [
                    str(entry.get("message"))
                    for entry in suggestions
                    if isinstance(entry, dict)
                    and isinstance(entry.get("message"), str)
                ]
                return "建議調整：" + "；".join(suggestion_messages) + "。"
            return "今天沒有可套用的配送參數建議。"
        deviation_details: list[str] = []
        for entries in (vehicle_items, zone_items):
            if not isinstance(entries, list):
                continue
            deviation_details.extend(
                str(entry.get("message"))
                for entry in entries
                if isinstance(entry, dict)
                and isinstance(entry.get("message"), str)
            )
        return (
            "今天回顧：" + " ".join(deviation_details)
            if deviation_details
            else "今天目前沒有記錄到配送偏差。"
        )
    if tool == "explain_unassigned":
        order_id = item.get("order_id")
        reason = item.get("reason")
        if item.get("status") == "ORDER_NOT_FOUND":
            return f"找不到訂單 {order_id}，資料中沒有這張訂單。"
        reason_labels = {
            "OVER_VEHICLE_CAPACITY": "這張單比最大的那台車還重，空車也裝不下",
            "CAPACITY_LIMIT": "每一台車的載重餘裕都不夠裝這張單",
            "SERVICE_ZONE_UNAVAILABLE": "沒有車負責這一區",
            "TIME_OR_ROUTE_CONFLICT": "配送時段排不下，或是繞過去會讓別的單遲到",
            "TIME_WINDOW_CONFLICT": "配送時段排不下",
            "VEHICLE_UNAVAILABLE": "今天可用的車不夠",
            "UNASSIGNABLE": "載重、責任區、配送時段三個條件湊不出可行的安排",
            "UNASSIGNED_BY_SOLVER": "載重、責任區、配送時段三個條件湊不出可行的安排",
        }
        if isinstance(order_id, str):
            if str(reason) == "ORDER_IS_ASSIGNED":
                return f"{order_id} 其實已經排進去了，不在未安排清單裡。"
            label = reason_labels.get(
                str(reason), str(reason) if reason else "目前的條件下排不進去"
            )
            return f"{order_id} 今天排不進去。\n原因：{label}。"
        return "目前沒有排不進去的訂單。"
    if tool == "explain_assignment":
        order_id = item.get("order_id")
        if item.get("status") == "ORDER_NOT_FOUND":
            return f"找不到訂單 {order_id}，資料中沒有這張訂單。"
        assignment_reason = item.get("assignment_reason")
        source_utterance = item.get("source_utterance")
        if isinstance(order_id, str) and isinstance(assignment_reason, str):
            suffix = (
                f" 原句：{source_utterance}"
                if isinstance(source_utterance, str) and source_utterance
                else ""
            )
            return assignment_reason + suffix
        vehicle_id = item.get("vehicle_id")
        load = _message_number(item.get("planned_load_kg"))
        if isinstance(order_id, str) and isinstance(vehicle_id, str):
            suffix = f"，該車計畫載重 {load} kg" if load is not None else ""
            return f"訂單 {order_id} 目前安排在 {vehicle_id}{suffix}；詳細依據可在方案明細查看。"
        if isinstance(order_id, str):
            return f"訂單 {order_id} 目前沒有安排到車輛。"
        return "目前沒有可說明的訂單分配。"
    if tool == "compare_strategies":
        strategy_names: list[str] = [
            str(entry.get("objective"))
            for entry in item.get("strategies", [])
            if isinstance(entry, dict) and isinstance(entry.get("objective"), str)
        ]
        name_text = "、".join(strategy_names) if strategy_names else "三種"
        return f"已完成 {name_text} 方案比較，請查看距離、時間與載重取捨。"
    if tool == "simulate_delay":
        delay = item.get("delay")
        if isinstance(delay, dict):
            minutes = _message_number(delay.get("delay_minutes"))
            affected = _message_number(delay.get("affected_order_count"))
            ids = _message_ids(delay.get("affected_orders"))
            slack = delay.get("tightest_slack_minutes")
            if minutes is not None and affected is not None:
                if ids:
                    return (
                        f"全隊晚 {minutes} 分鐘的話，有 {affected} 張會掉出配送時段：\n{ids}"
                    )
                spare = (
                    f"最緊的一張還有 {round(slack)} 分鐘裕度，"
                    if isinstance(slack, (int, float))
                    else ""
                )
                return (
                    f"全隊晚 {minutes} 分鐘還撐得住，沒有訂單會掉出配送時段。\n"
                    f"{spare}所以這個延誤吸收得掉。"
                )
        return "已完成配送延遲風險試算，請查看受影響訂單。"
    if tool == "change_vehicle_availability":
        vehicle_id = item.get("affected_vehicle_id", item.get("vehicle_id"))
        plan = item.get("plan")
        if isinstance(vehicle_id, str) and isinstance(plan, dict):
            assigned = _message_number(plan.get("assigned_order_count"))
            unassigned_orders = plan.get("unassigned_orders")
            if assigned is not None and isinstance(unassigned_orders, list):
                action = (
                    "今天停駛"
                    if item.get("requested_status") == "UNAVAILABLE"
                    else "恢復出車"
                )
                return (
                    f"{vehicle_label(vehicle_id)}{action}試算完成：目前可安排 {assigned} 張，"
                    f"未安排 {len(unassigned_orders)} 張"
                    + (
                        f"。{item['conflict_summary']}。"
                        if isinstance(item.get("conflict_summary"), str)
                        and item.get("conflict_summary")
                        else "；"
                    )
                    + "請檢查後再確認。"
                )
        return f"{vehicle_label(vehicle_id)}的出勤狀態試算未完成，請檢查車輛編號。"
    if tool == "change_order_constraint":
        order_id = item.get("order_id")
        slot = item.get("time_slot")
        if item.get("status") == "MISSING_ORDER_ID":
            message = item.get("message")
            if isinstance(message, str) and message.strip():
                return message
        return f"已試算訂單 {order_id} 改為 {slot} 時段，方案尚未套用；請查看新的方案卡。"
    if tool == "change_frozen_stops":
        stop_count = _message_number(item.get("selected_stop_count"))
        action = "凍結" if item.get("action") == "FREEZE" else "解除凍結"
        return f"已{action} {stop_count or '指定'} 個站點，對話中的方案卡尚未套用。"
    if tool == "reassign_order_preview":
        order_id = item.get("order_id")
        target = item.get("target_vehicle_id")
        if item.get("status") in {"NOT_FOUND", "ORDER_NOT_FOUND"}:
            return f"找不到訂單 {order_id}，這次換車／改派沒有執行，原方案沒有變更。"
        if item.get("status") == "MISSING_TARGET_VEHICLE":
            # 沒有人講車的時候就問, 不要挑一台來湊。
            return f"{order_id} 要換到哪一台車？講一台給我，現在的方案沒有變更。"
        return f"已試算訂單 {order_id} 改派至 {vehicle_label(target)} 的方案，請檢查方案卡。"
    if tool == "prioritize_order_preview":
        order_id = item.get("order_id")
        if item.get("status") in {"NOT_FOUND", "ORDER_NOT_FOUND"}:
            return f"找不到訂單 {order_id}，無法提前或先送，原方案沒有變更。"
        eta = _eta_clock(item.get("estimated_eta"))
        suffix = f"，預估 {eta} 到達" if eta else ""
        return f"已試算訂單 {order_id} 提前配送{suffix}；對話中的新方案卡尚未套用。"
    if tool == "remove_order_preview":
        order_id = item.get("order_id")
        assigned = _message_number(item.get("assigned_order_count"))
        if isinstance(order_id, str) and assigned is not None:
            return (
                f"已試算今天不配送 {order_id}：目前方案將安排 "
                f"{assigned} 張訂單；原方案尚未變更。"
            )
        return f"找不到訂單 {order_id}，目前沒有移除方案可預覽。"
    if tool == "enforce_hard_time_windows":
        return "已重算成每一張都守住配送時段，下面的方案卡有代價。還沒套用。"
    if tool == "query_plan_version":
        plan_id = item.get("plan_id")
        version = _message_number(item.get("version"))
        if isinstance(plan_id, str) and version is not None:
            return f"目前方案版本是 {plan_id} 第 {version} 版。"
        return "目前還沒有可查詢的方案版本。"
    if tool in {
        "preview_urgent_insert",
        "preview_structured_urgent_insert",
        "preview_multiple_urgent_insert",
    }:
        order_ids = _message_ids(item.get("order_ids"))
        if item.get("status") == "ORDER_ID_EXISTS":
            return (
                f"{order_ids or '這個編號'} 今天已經有了，不能用同一個編號再插一次。\n"
                "換一個編號再給我一次。"
            )
        if item.get("status") == "UNASSIGNABLE":
            return (
                f"{order_ids or '這批急單'} 插不進去。\n"
                "所有車在載重、責任區或配送時段上都湊不出可行的位置，原方案沒有變更。"
            )
        if item.get("status") == "URGENT_ORDER_INVALID":
            # Saying 「已經算好」 over a rejected preview is the worst kind of
            # wrong: the card never appears and the dispatcher waits for it.
            validation = item.get("validation")
            reasons = dict.fromkeys(
                str(error["message"])
                for error in (validation or {}).get("errors", [])
                if isinstance(error, dict) and error.get("message")
            )
            detail = "\n".join(reasons)
            return (
                f"{order_ids or '這批急單'} 沒有算成，資料沒通過檢查：\n{detail}\n"
                "原方案沒有變更。把上面這幾筆改好再給我一次。"
                if detail
                else f"{order_ids or '這批急單'} 的資料沒通過檢查，原方案沒有變更。"
            )
        # This path fires when the model previews straight from a complete
        # description without the review step. It produces a whole-plan diff,
        # not the pick-one option cards, so the sentence has to carry the
        # numbers itself; pointing at cards that never render leaves the
        # dispatcher waiting for something that is not coming.
        return _multiple_urgent_preview_sentence(item, order_ids)
    return "這項操作已完成，請查看方案明細。"


def _scope_chat_message(evidence: list[dict[str, Any]], fallback: str) -> str:
    """Render only deterministic tool evidence as the operator-facing reply."""
    item = evidence[-1] if evidence else {}
    fallback_message = _json_message(fallback)
    raw_message = _json_message(item.get("message"))
    message = raw_message or (
        item.get("message") if isinstance(item.get("message"), str) else None
    )
    if (
        isinstance(message, str)
        and message.strip()
        and not message.startswith("已完成確定性工具計算")
    ):
        return message.strip()
    if fallback_message and not fallback_message.startswith("已完成確定性工具計算"):
        return fallback_message
    if item:
        return _operator_tool_template(item.get("tool"), item)
    return _json_message(fallback) or fallback


def _modification_option_label(kind: str) -> tuple[str, str]:
    labels = {
        "REASSIGN_VEHICLE": ("調整車輛", "依指定車輛重新安排"),
        "TIME_SLOT_CHANGE": ("調整時段", "依新配送時段重新安排"),
        "REMOVE_ORDER": ("改明天送", "從今日方案移除"),
        "PRIORITIZE_ORDER": ("先送這單", "提高該車剩餘站點順序"),
        "FREEZE_VEHICLE": ("凍結車線", "保留指定車輛目前站點"),
        "FREEZE_STOPS": ("凍結站點", "保留指定站點目前位置"),
        "HARD_TIME_WINDOWS": ("硬性時段", "所有配送時段維持硬性限制"),
    }
    return labels.get(kind, ("修改方案", "依目前要求重新安排"))


def _persist_plan_change_preview(
    context: Any,
    base_record: PlanRecord,
    dataset_record: DatasetRecord,
) -> dict[str, Any] | None:
    """Persist one deterministic Agent candidate as an immutable card preview."""
    candidate = getattr(context, "pending_preview_plan", None)
    kind = getattr(context, "pending_preview_kind", None)
    if candidate is None or not isinstance(kind, str):
        return None
    changed_dataset = getattr(context, "pending_preview_dataset", None) or dataset_record.dataset
    preview_matrix = (
        SimulatedRouteProvider().build(changed_dataset)
        if dataset_hash(changed_dataset) != dataset_hash(dataset_record.dataset)
        else base_record.matrix
    )
    validation = validate_plan(changed_dataset, candidate, preview_matrix)
    baseline_unassigned = set(base_record.plan.unassigned_orders)
    candidate_unassigned = set(candidate.unassigned_orders)
    if not validation.valid or not candidate_unassigned.issubset(baseline_unassigned):
        return None

    preview_dataset_id = base_record.dataset_id
    if dataset_hash(changed_dataset) != dataset_hash(dataset_record.dataset):
        preview_dataset_id = f"DS-{uuid4().hex[:12].upper()}"
        changed_report = validate_dataset(changed_dataset)
        changed_record = DatasetRecord(
            dataset_id=preview_dataset_id,
            dataset=changed_dataset,
            validation=changed_report,
            matrix=preview_matrix,
            created_at=datetime.now(UTC).isoformat(),
        )
        store.add_dataset(changed_record)
        repository.save_dataset(
            preview_dataset_id,
            changed_dataset,
            changed_report,
            preview_matrix,
            changed_record.created_at,
        )

    preview_version = max(store.plans.get(base_record.plan_id, {0: None})) + 1
    preview_state = (
        base_record.state if base_record.state in {"LOADED", "DISPATCHED"} else "PROPOSED"
    )
    preview_plan = candidate.model_copy(update={"state": preview_state})
    preview_record = PlanRecord(
        plan_id=base_record.plan_id,
        dataset_id=preview_dataset_id,
        version=preview_version,
        state=preview_state,
        plan=preview_plan,
        validation=validation,
        matrix=preview_matrix,
        created_at=datetime.now(UTC).isoformat(),
    )
    store.add_plan(preview_record, make_current=False)
    repository.save_plan(preview_record, make_current=False)
    diff = compute_plan_diff(base_record.plan, preview_plan)
    title, rationale = _modification_option_label(kind)
    order_id = getattr(context, "pending_preview_order_id", None)
    estimated_eta = next(
        (
            stop.eta
            for route in preview_plan.routes
            for stop in route.stops
            if stop.order_id == order_id
        ),
        None,
    )
    if estimated_eta is None:
        estimated_eta = next(
            (
                stop.eta
                for route in base_record.plan.routes
                for stop in route.stops
                if stop.order_id == order_id
            ),
            None,
        )
    slack_values = [route.max_load_kg - route.planned_load_kg for route in preview_plan.routes]
    # A modification card describes the cost of choosing this candidate.  The
    # plan diff remains the complete audit record, but a shorter route or ETA
    # must not be rendered as a negative insertion cost to the dispatcher.
    distance_cost_m = max(0, diff["total_distance_delta_m"])
    duration_cost_s = max(0, diff["total_duration_delta_s"])
    priority_metadata = getattr(context, "pending_preview_metadata", None)
    card_rationale = rationale
    if kind == "PRIORITIZE_ORDER" and isinstance(priority_metadata, dict):
        metadata_rationale = priority_metadata.get("rationale")
        if isinstance(metadata_rationale, str):
            card_rationale = metadata_rationale
    return {
        "option_id": f"CHANGE-{preview_version}",
        "label": "方案 A",
        "title": title,
        "rationale": card_rationale,
        "mode": "MODIFICATION",
        "change": {"kind": kind, "order_id": order_id},
        "plan_id": base_record.plan_id,
        "base_version": base_record.version,
        "preview_version": preview_version,
        "feasible": True,
        "selectable": True,
        "requires_human_confirmation": True,
        "inserted_orders": [],
        "estimated_eta": estimated_eta,
        "cost": {
            "distance_delta_m": distance_cost_m,
            "distance_delta_km": round(distance_cost_m / 1000, 1),
            "duration_delta_s": duration_cost_s,
            "duration_delta_min": round(duration_cost_s / 60, 1),
            "vehicle_change_count": len(diff["reassigned_orders"]),
            "minimum_capacity_slack_kg": round(min(slack_values, default=0.0), 1),
        },
        "affected_vehicle_count": len(
            {
                item["vehicle_id"]
                for item in diff["vehicle_load_changes"]
                if item["delta_load_kg"] != 0
            }
            | {
                item["from_vehicle_id"]
                for item in diff["reassigned_orders"]
            }
            | {
                item["to_vehicle_id"]
                for item in diff["reassigned_orders"]
            }
        ),
        "moved_order_count": len(diff["sequence_changes"]),
        "before": _plan_payload(base_record)["summary"],
        "after": _plan_payload(preview_record)["summary"],
        "validator": validation.model_dump(mode="json"),
        "diff": diff,
        "current_state": (
            priority_metadata.get("current_state")
            if isinstance(priority_metadata, dict)
            else None
        ),
        "replanned_route": (
            priority_metadata.get("replanned_order_ids")
            if isinstance(priority_metadata, dict)
            else None
        ),
        "sacrificed_order_ids": (
            priority_metadata.get("sacrificed_order_ids", [])
            if isinstance(priority_metadata, dict)
            else []
        ),
    }


def _save_agent_session(session_id: str, session: AgentSession) -> None:
    if len(session.history) > 12:
        session.history = session.history[-12:]
    repository.save_agent_session(
        session_id,
        _session_payload(session),
        datetime.now(UTC).isoformat(),
    )


def _clear_agent_planning_context(session: AgentSession) -> None:
    """Drop lifecycle state that belongs to a previous dataset or plan."""
    session.dataset_id = None
    session.plan_id = None
    session.plan_version = None
    session.order_id = None
    session.vehicle_id = None
    session.strategy = None
    session.rule_source_utterance = None
    session.frozen_stop_count = 0
    session.frozen_stop_ids = ()
    session.pending_fields = ()
    session.last_preview_version = None
    session.last_tool = None
    session.pending_order = None
    session.urgent_workflow = UrgentWorkflowState().model_dump(mode="json")
    session.history = []


def _sync_direct_urgent_preview_session(
    request: Request,
    plan_id: str,
    base_plan_version: int,
    bundles: list[UrgentOrderBundleRequest],
    response: Any,
) -> None:
    """Keep a browser demo preview and its following chat turn on one candidate."""
    session_id = request.headers.get("X-Dispatch-Session")
    if not session_id or not isinstance(response, dict) or response.get("feasible") is not True:
        return
    preview_version = response.get("preview_version")
    if not isinstance(preview_version, int):
        return
    base_record = store.get_plan(plan_id, base_plan_version)
    if base_record is None:
        return
    drafts: list[UrgentOrderDraft] = []
    for bundle in bundles:
        weights = {package.weight_kg for package in bundle.packages}
        per_package_weight = next(iter(weights)) if len(weights) == 1 else None
        drafts.append(
            UrgentOrderDraft.model_validate(
                bundle.order.model_dump(mode="json", exclude={"note"})
                | {"package_weight_kg": per_package_weight}
            )
        )
    session = agent_sessions.get(session_id)
    if session is None:
        persisted_session = repository.load_agent_session(session_id)
        session = _session_from_payload(persisted_session) if persisted_session else AgentSession()
        agent_sessions[session_id] = session
    session.dataset_id = base_record.dataset_id
    session.plan_id = plan_id
    session.plan_version = base_plan_version
    session.order_id = drafts[-1].order_id if drafts else None
    session.urgent_workflow = UrgentWorkflowState(
        stage="PREVIEW_READY", orders=drafts
    ).model_dump(mode="json")
    session.last_preview_version = preview_version
    session.last_tool = "urgent_insertion_workflow"
    _save_agent_session(session_id, session)


@app.post("/api/v1/agent/chat")
async def agent_chat(payload: ChatRequest, request: Request) -> Any:
    if not settings.openai_api_key:
        return _error(
            request, 503, "AGENT_UNAVAILABLE", "OpenAI 憑證未設定; 確定性 REST 功能仍可使用."
        )
    session = agent_sessions.get(payload.session_id)
    if session is None:
        persisted_session = repository.load_agent_session(payload.session_id)
        session = _session_from_payload(persisted_session) if persisted_session else AgentSession()
        agent_sessions[payload.session_id] = session
    explicit_plan_id = payload.context.get("plan_id")
    explicit_dataset_id = payload.context.get("dataset_id")
    plan_changed = (
        isinstance(explicit_plan_id, str)
        and session.plan_id not in {None, explicit_plan_id}
    )
    dataset_changed = (
        isinstance(explicit_dataset_id, str)
        and session.dataset_id not in {None, explicit_dataset_id}
    )
    if plan_changed or dataset_changed:
        _clear_agent_planning_context(session)
        _save_agent_session(payload.session_id, session)
    # An explicit dataset pointer starts a new planning context; do not let a
    # stale persisted plan from the same conversation override it.  When the
    # caller omits both pointers, resume the last structured plan pointer.
    context_plan_id = (
        explicit_plan_id
        if isinstance(explicit_plan_id, str)
        else None
        if isinstance(explicit_dataset_id, str)
        else session.plan_id
    )
    context_dataset_id = explicit_dataset_id or session.dataset_id
    if not isinstance(context_dataset_id, str):
        context_dataset_id = None
    context_order_id = payload.context.get("order_id")
    if not isinstance(context_order_id, str):
        context_order_id = session.order_id
    context_vehicle_id = payload.context.get("vehicle_id")
    if not isinstance(context_vehicle_id, str):
        context_vehicle_id = session.vehicle_id
    record: PlanRecord | None = None
    dataset_record: DatasetRecord | None = None
    if isinstance(context_plan_id, str):
        plan_version = payload.context.get("plan_version")
        if not isinstance(plan_version, int):
            plan_version = None
        record = store.get_plan(context_plan_id, plan_version)
        if record is None:
            return _error(request, 404, "PLAN_NOT_FOUND", "找不到說明所需的規劃版本。")
        dataset_record = store.get_dataset(record.dataset_id)
        if dataset_record is None:
            return _error(request, 404, "DATASET_NOT_FOUND", "找不到說明所需的資料集。")
    elif context_dataset_id is not None:
        dataset_record = store.get_dataset(context_dataset_id)
        if dataset_record is None:
            return _error(request, 404, "DATASET_NOT_FOUND", "找不到目前資料集。")
    if record is not None and dataset_record is not None:
        # Continuations always use the immutable matrix attached to the
        # selected plan so explanations and previews cannot drift.
        dataset, matrix = dataset_record.dataset, record.matrix
    elif dataset_record is not None:
        # A new plan requested through chat must resolve its provider matrix
        # before Runner.run.  The selected deterministic planning tool then
        # receives exactly this matrix; it must not silently rebuild a
        # different source.  Missing credentials remain an explicit
        # simulated warning, while provider HTTP failures are surfaced.
        try:
            matrix = _build_matrix(dataset_record.dataset, prefer_live=True)
        except GoogleRoutesProviderError as exc:
            return _error(
                request,
                502,
                "PROVIDER_UNAVAILABLE",
                "Google Routes 即時矩陣無法取得，暫時無法建立配送方案。",
                provider="GOOGLE",
                operation="computeRouteMatrix",
                provider_error=exc.code,
                provider_error_category=exc.category,
                fallback_used=False,
                retryable=exc.code in {"GOOGLE_TIMEOUT", "GOOGLE_REQUEST_FAILED"},
            )
        dataset = dataset_record.dataset
    else:
        dataset, matrix = _empty_agent_dataset()
    if (
        dataset_record is not None
        and session.dataset_id is not None
        and session.dataset_id != dataset_record.dataset_id
        and not dataset_changed
    ):
        _clear_agent_planning_context(session)
        _save_agent_session(payload.session_id, session)
    if dataset_record is not None:
        session.dataset_id = dataset_record.dataset_id

    if isinstance(payload.context.get("order_id"), str):
        session.order_id = payload.context["order_id"]
    session_vehicle_id = payload.context.get("vehicle_id")
    if isinstance(session_vehicle_id, str):
        context_vehicle_id = session_vehicle_id
        session.vehicle_id = session_vehicle_id
    if isinstance(payload.context.get("strategy"), str):
        session.strategy = str(payload.context["strategy"])
    if isinstance(payload.context.get("frozen_stop_count"), int):
        session.frozen_stop_count = int(payload.context["frozen_stop_count"])
    raw_frozen_stop_ids = payload.context.get("frozen_stop_ids")
    if isinstance(raw_frozen_stop_ids, list):
        session.frozen_stop_ids = tuple(
            item for item in raw_frozen_stop_ids if isinstance(item, str)
        )
    if record is not None:
        session.plan_id = record.plan_id
        session.plan_version = record.version
    current_stage = stage_for_state(record.state) if record is not None else "PRE_LOAD"
    raw_timeline_minutes = payload.context.get("timeline_minutes")
    timeline_minutes = (
        raw_timeline_minutes
        if isinstance(raw_timeline_minutes, int) and not isinstance(raw_timeline_minutes, bool)
        else None
    )
    if record is not None and current_stage == "DISPATCHED":
        session.frozen_stop_ids = tuple(
            item["order_id"]
            for route in route_progress(record.plan, timeline_minutes)
            for item in route["stop_statuses"]
            if item["status"] == "COMPLETED"
        )
        session.frozen_stop_count = len(session.frozen_stop_ids)

    urgent_state = UrgentWorkflowState.model_validate(session.urgent_workflow or {})

    # A visible urgent card is itself an immutable candidate plan.  Let a
    # subsequent plan-change tool operate on the selected preview candidate so
    # that "send this order first" can reorder the inserted stop instead of
    # looking for that not-yet-confirmed order in the original dataset.
    agent_record = record
    agent_dataset_record = dataset_record
    agent_dataset = dataset
    agent_matrix = matrix
    preview_record_for_change: PlanRecord | None = None
    if (
        urgent_state.stage == "PREVIEW_READY"
        and record is not None
        and session.last_preview_version is not None
    ):
        preview_record_for_change = store.get_plan(
            record.plan_id, session.last_preview_version
        )
        if preview_record_for_change is not None:
            preview_dataset_record = store.get_dataset(preview_record_for_change.dataset_id)
            if preview_dataset_record is not None:
                agent_record = preview_record_for_change
                agent_dataset_record = preview_dataset_record
                agent_dataset = preview_dataset_record.dataset
                agent_matrix = preview_record_for_change.matrix

    def finish_urgent_workflow(
        understanding: UrgentUnderstanding,
        runner_result: Any,
        *,
        entry_tool: str = "urgent_insertion_workflow",
    ) -> Any:
        understanding = _normalize_urgent_order_context(understanding, dataset)
        existing_ids = {order.order_id for order in dataset.orders}
        workflow = advance_urgent_workflow(
            urgent_state,
            understanding,
            fixture_lookup=get_demo_urgent_order,
            existing_order_ids=existing_ids,
        )
        orders = [item.model_dump(mode="json") for item in workflow.state.orders]
        missing = [item.model_dump(mode="json") for item in workflow.missing_by_order]
        preview: dict[str, Any] | None = None
        next_state = workflow.state
        if workflow.should_preview:
            if record is None or dataset_record is None:
                return _error(
                    request,
                    409,
                    "PLAN_REQUIRED",
                    "請先建立今天的配送方案，再產生插單預覽。",
                )
            bundles = []
            for draft in workflow.complete_orders:
                order = draft.to_order()
                bundles.append(
                    UrgentOrderBundleRequest(
                        order=UrgentOrderRequest.model_validate(
                            order.model_dump(exclude={"packages", "total_weight_kg"})
                        ),
                        packages=list(order.packages),
                    )
                )
            preview_response = _urgent_insert_preview_response(
                record.plan_id,
                record.version,
                bundles,
                request,
                include_unassignable_option=True,
            )
            if isinstance(preview_response, JSONResponse):
                # A validation error must not discard the draft. Keep the
                # workflow in REVIEW_READY so the next turn can supply a
                # corrected field instead of being routed as a new request.
                validation_errors = _urgent_validation_errors(preview_response)
                pending_order_id = _urgent_validation_order_id(preview_response, orders)
                next_state = workflow.state.model_copy(
                    update={
                        "stage": "REVIEW_READY",
                        "pending_order_id": pending_order_id,
                    }
                )
                session.urgent_workflow = next_state.model_dump(mode="json")
                session.pending_fields = tuple(
                    f"{item['order_ref']}.{field}"
                    for item in missing
                    for field in item["missing_fields"]
                ) or tuple(
                    str(error["path"])
                    for error in validation_errors
                    if isinstance(error.get("path"), str)
                )
                if orders and isinstance(orders[-1].get("order_id"), str):
                    session.order_id = str(orders[-1]["order_id"])
                session.last_tool = "urgent_insertion_workflow"
                session.history.extend(
                    [
                        ("user", _safe_session_text(payload.message)),
                        ("assistant", "插單資料未通過驗證，請修正欄位後再試。"),
                    ]
                )
                _save_agent_session(payload.session_id, session)
                return preview_response
            preview = preview_response
            # 只留剛剛真的算進去的那幾張。前面「加一張到信義區」留下的半張單
            # 還掛在草稿裡時, 下一句「用 A, 但這單先送」會被讀成「跳過那張的
            # 檢查」, 回一句「不能跳過資料檢查」——調度員問的是提前送。
            next_state = workflow.state.model_copy(
                update={
                    "stage": "PREVIEW_READY",
                    "orders": list(workflow.complete_orders),
                }
            )
            session.last_preview_version = int(preview["preview_version"])
        # A cancelled draft is nothing in progress. Persisting the CANCELLED
        # stage kept the urgent tools switched off, so the very next 「再加一張
        # 急單」 had nowhere to go: the dispatcher dropped one order and could
        # not start another. This turn still reports CANCELLED so the reply
        # says 「已取消這次臨時插單」; only what is carried forward resets.
        session.urgent_workflow = (
            UrgentWorkflowState().model_dump(mode="json")
            if next_state.stage == "CANCELLED"
            else next_state.model_dump(mode="json")
        )
        session.pending_fields = tuple(
            f"{item['order_ref']}.{field}"
            for item in missing
            for field in item["missing_fields"]
        )
        if next_state.stage == "CANCELLED":
            session.pending_fields = ()
            session.order_id = None
        elif orders and isinstance(orders[-1].get("order_id"), str):
            session.order_id = str(orders[-1]["order_id"])
        evidence_data: dict[str, Any] = {
            "stage": next_state.stage,
            "orders": orders,
            "pending_order_id": next_state.pending_order_id,
            "missing_by_order": missing,
            "duplicate_order_ids": workflow.duplicate_order_ids,
            "requires_human_confirmation": preview is not None,
        }
        if preview is not None:
            evidence_data["preview"] = preview
            evidence_data["options"] = preview.get("options", [])
        final_output = _urgent_workflow_message(
            next_state.stage,
            orders,
            missing,
            workflow.duplicate_order_ids,
            preview,
        )
        session.history.extend(
            [
                ("user", _safe_session_text(payload.message)),
                ("assistant", _safe_session_text(final_output)),
            ]
        )
        session.last_tool = "urgent_insertion_workflow"
        _save_agent_session(payload.session_id, session)
        provider_runtime_state["openai"] = "connected"
        urgent_usage = getattr(
            getattr(runner_result, "context_wrapper", None), "usage", None
        )
        # The public HTTP contract exposes the completed deterministic workflow
        # as the evidence record. The semantic entry tool is an internal
        # routing detail; returning it here would make callers inspect a
        # partial draft instead of the final workflow state.
        workflow_evidence = [
            {"tool": "urgent_insertion_workflow", "data": evidence_data}
        ]
        return {
            "session_id": payload.session_id,
            "agent_run_id": f"RUN-{uuid4().hex[:12].upper()}",
            "message": final_output,
            "evidence": workflow_evidence,
            "requires_human_confirmation": preview is not None,
            "usage": {
                "total_tokens": int(getattr(urgent_usage, "total_tokens", 0) or 0),
                "tool_calls": 0,
            },
            "provider_mode": record.matrix.provider_mode if record else matrix.provider_mode,
            "plan_id": record.plan_id if record else None,
            "plan_version": record.version if record else None,
            "runner_result_type": type(runner_result).__name__,
            "request_id": _request_id(request),
        }

    # Once an urgent preview card is visible, the next turn belongs to the
    # plan-change Agent tools.  The dedicated urgent interpreter must not
    # reinterpret a card choice such as "use A and send this order first" as
    # another draft update.  New urgent orders and cancellation remain
    # available through the strict begin_urgent_insertion tool below.
    plan_change_turn = urgent_state.stage == "PREVIEW_READY"
    if payload.action == "PREVIEW_URGENT" and urgent_state.stage == "REVIEW_READY":
        urgent_understanding = UrgentUnderstanding(
            is_urgent_insertion=True,
            action="PREVIEW",
        )
        urgent_run = None
    elif session.last_tool == "preview_dispatch_rule" and urgent_state.stage == "IDLE":
        # A rule clarification is an Agent-owned continuation.  The urgent
        # interpreter is deliberately skipped for this lifecycle edge so a
        # numeric rule follow-up cannot be mistaken for a new urgent draft.
        urgent_understanding = UrgentUnderstanding(is_urgent_insertion=False, action="NONE")
        urgent_run = None
    elif urgent_state.stage == "IDLE":
        # A fresh conversation must reach the main dispatch Agent so its
        # strict tool descriptions decide between urgent-intake entry points
        # and ordinary plan operations.  The structured urgent interpreter is
        # reserved for merging an already active draft or handling a visible
        # preview; it must not act as a pre-agent router for a new message.
        urgent_understanding = UrgentUnderstanding(is_urgent_insertion=False, action="NONE")
        urgent_run = None
    else:
        try:
            existing_order_ids = {order.order_id for order in dataset.orders}
            vehicle_choices = [
                {"vehicle_id": vehicle.vehicle_id, "vehicle_name": vehicle.vehicle_name}
                for vehicle in dataset.vehicles
            ]
            understand_parameters = inspect.signature(understand_urgent_message).parameters
            understand = cast(Any, understand_urgent_message)
            if "vehicle_choices" in understand_parameters:
                urgent_understanding, urgent_run = await understand(
                    payload.message,
                    urgent_state,
                    vehicle_choices=vehicle_choices,
                    **(
                        {"existing_order_ids": sorted(existing_order_ids)}
                        if "existing_order_ids" in understand_parameters
                        else {}
                    ),
                )
            else:
                urgent_understanding, urgent_run = await understand(
                    payload.message,
                    urgent_state,
                )
        except InputGuardrailTripwireTriggered:
            return _error(
                request,
                400,
                "PROMPT_INJECTION_BLOCKED",
                "訊息包含不可執行的規則繞過要求。",
            )
        except Exception as exc:
            provider_runtime_state["openai"] = "failed"
            status_code, error_code, message, retryable = _classify_agent_error(exc)
            return _error(
                request,
                status_code,
                error_code,
                message,
                provider="OPENAI",
                exception_type=type(exc).__name__,
                fallback_used=False,
                retryable=retryable,
            )

    if plan_change_turn and urgent_understanding.is_urgent_insertion:
        # A plan-change turn may still introduce a genuinely new urgent draft,
        # but a bare reference to the already shown card must stay with the
        # main Agent. Use only the interpreter's strict provenance fields for
        # this boundary; never inspect the user's text here.
        has_new_supplied_field = any(
            bool(order.supplied_fields) for order in urgent_understanding.orders
        )
        is_draft_cancellation = urgent_understanding.action == "CANCEL"
        if not has_new_supplied_field and not is_draft_cancellation:
            urgent_understanding = UrgentUnderstanding(
                is_urgent_insertion=False,
                action="NONE",
            )
            urgent_run = None

    # While a draft is collecting fields, only a structured field update or an
    # explicit cancellation belongs to that workflow. Questions about the
    # current plan (for example, whether it would overload a vehicle) must go
    # back to the main Agent instead of replaying the draft's missing fields.
    # This boundary uses only the strict typed action/provenance result, never
    # text matching; an update that actually supplies fields remains in intake.
    active_draft_update = (
        urgent_understanding.action in {"ADD_OR_UPDATE", "MODIFY"}
        and any(order.supplied_fields for order in urgent_understanding.orders)
    )
    active_draft_cancel = urgent_understanding.action == "CANCEL"
    if urgent_state.stage == "COLLECTING" and not (
        active_draft_update or active_draft_cancel
    ):
        urgent_understanding = UrgentUnderstanding(
            is_urgent_insertion=False,
            action="NONE",
        )
        urgent_run = None

    if urgent_understanding.is_urgent_insertion or urgent_understanding.action != "NONE":
        return finish_urgent_workflow(urgent_understanding, urgent_run)

    # Context identifiers are application-controlled data. Include only the
    # selected order identifier as a hint so the model must still invoke the
    # allowlisted deterministic tool instead of receiving precomputed facts.
    # Never replay an earlier action request as part of the current instruction.
    # Multi-turn references are resolved from the structured session pointers
    # below (plan/version/order/vehicle/last tool/pending fields). This prevents
    # an earlier incident or urgent order from competing with the current turn.
    agent_message = payload.message
    vehicle_metadata: list[dict[str, Any]] = []
    for vehicle in sorted(agent_dataset.vehicles, key=lambda item: item.vehicle_id):
        route = next(
            (
                item
                for item in (agent_record.plan.routes if agent_record is not None else [])
                if item.vehicle_id == vehicle.vehicle_id
            ),
            None,
        )
        vehicle_metadata.append(
            {
                "vehicle_id": vehicle.vehicle_id,
                "vehicle_name": vehicle.vehicle_name,
                "route_distance_km": round((route.total_distance_m if route else 0) / 1000, 1),
                "stop_count": len(route.order_ids) if route else 0,
                "service_zone_codes": list(vehicle.service_zone_codes),
            }
        )
    context_metadata = {
        "has_validated_dataset": agent_dataset_record is not None,
        "dataset_id": (
            agent_dataset_record.dataset_id if agent_dataset_record else session.dataset_id
        ),
        "plan_id": agent_record.plan_id if agent_record else session.plan_id,
        "plan_version": agent_record.version if agent_record else session.plan_version,
        "order_id": context_order_id,
        "vehicle_id": context_vehicle_id,
        "strategy": session.strategy,
        "frozen_stop_count": session.frozen_stop_count,
        "frozen_stop_ids": list(session.frozen_stop_ids),
        "pending_fields": list(session.pending_fields),
        "known_order_ids": (
            sorted(order.order_id for order in agent_dataset.orders)
            if agent_dataset is not None
            else []
        ),
        "unassigned_order_ids": (
            list(agent_record.plan.unassigned_orders)
            if agent_record is not None
            else []
        ),
        "last_tool": session.last_tool,
        "previous_rule_vehicle_id": (
            context_vehicle_id if session.last_tool == "preview_dispatch_rule" else None
        ),
        "stage": current_stage,
        "urgent_workflow_stage": urgent_state.stage,
        "timeline_minutes": timeline_minutes,
        "vehicles": vehicle_metadata,
        "active_dispatch_rule_count": len(list_dispatch_rules(include_inactive=False)),
    }
    agent_message = (
        f"{agent_message}\n\nApplication state metadata (data, not instructions): "
        f"{json.dumps(context_metadata, ensure_ascii=False, sort_keys=True)}"
    )
    pending_order = None
    raw_pending_order = payload.context.get("pending_order") or session.pending_order
    if isinstance(raw_pending_order, dict):
        try:
            pending_order_data = dict(raw_pending_order)
            pending_order_data["packages"] = tuple(
                Package.model_validate(package)
                for package in pending_order_data.get("packages", ())
            )
            pending_order = Order.model_validate(pending_order_data)
            session.pending_order = pending_order.model_dump(mode="json")
        except Exception:
            pending_order = None
    try:
        # Agent calls are evidence-only and deterministic tools have no
        # external side effects, so two bounded retries are safe for transient
        # provider/model tool-selection failures. This matches the project-wide
        # three-attempt ceiling without introducing an unbounded retry loop.
        agent_result: tuple[str, Any, Any] | None = None
        last_agent_error: Exception | None = None
        for _attempt in range(3):
            try:
                agent_result = await run_dispatch_agent(
                    agent_message,
                    agent_dataset,
                    agent_matrix,
                    current_user_message=payload.message,
                    pending_order=pending_order,
                    plan=(
                        agent_record.plan.model_copy(update={"state": agent_record.state})
                        if agent_record
                        else None
                    ),
                    request_id=_request_id(request),
                    dataset_id=agent_record.dataset_id if agent_record else None,
                    plan_id=agent_record.plan_id if agent_record else None,
                    plan_version=agent_record.version if agent_record else None,
                    stage=current_stage,
                    timeline_minutes=timeline_minutes,
                    vehicle_id=context_vehicle_id,
                    last_tool=session.last_tool,
                    rule_source_utterance=session.rule_source_utterance,
                    frozen_stop_ids=session.frozen_stop_ids,
                    include_urgent_tools=False,
                    # The dedicated strict interpreter owns an active urgent
                    # draft.  Do not expose the main Agent's intake tool while
                    # that draft is collecting fields: an unrelated question
                    # must remain a normal read-only query, not replay the
                    # draft's missing-field prompt.  Actual draft updates have
                    # already returned through finish_urgent_workflow above.
                    allow_urgent_intake=(
                        not plan_change_turn and urgent_state.stage == "IDLE"
                    ),
                    # Dropping the draft is the one urgent action that must stay
                    # reachable while a draft is open. Without it the only tool
                    # left that fits 「算了不要了」 is the refusal, which answers
                    # a question the dispatcher did not ask. It carries no order
                    # fields, so it cannot replay the missing-field prompt.
                    allow_urgent_cancel=(
                        not plan_change_turn and urgent_state.stage != "IDLE"
                    ),
                    plan_change_mode=plan_change_turn,
                )
                break
            except InputGuardrailTripwireTriggered:
                raise
            except Exception as exc:
                last_agent_error = exc
        if agent_result is None:
            assert last_agent_error is not None
            raise last_agent_error
        final_output, context, result = agent_result
        final_output = _scope_chat_message(context.evidence, final_output)
    except InputGuardrailTripwireTriggered:
        return _error(
            request,
            400,
            "PROMPT_INJECTION_BLOCKED",
            "訊息包含不可執行的規則繞過要求。",
        )
    except Exception as exc:
        # Do not serialize provider requests, headers, keys, or SDK internals.
        provider_runtime_state["openai"] = "failed"
        status_code, error_code, message, retryable = _classify_agent_error(exc)
        return _error(
            request,
            status_code,
            error_code,
            message,
            provider="OPENAI",
            exception_type=type(exc).__name__,
            fallback_used=False,
            retryable=retryable,
        )

    provider_runtime_state["openai"] = "connected"

    # The dedicated structured interpreter is the primary urgent-order gate.
    # If that model pass ever says NONE but the main Runner semantically selects
    # the safe intake tool, hand the same structured facts to the deterministic
    # workflow. This prevents an urgent request from being misapplied as a new
    # full plan without adding regex or keyword routing.
    urgent_intake = next(
        (item for item in context.evidence if item.get("tool") == "begin_urgent_insertion"),
        None,
    )
    if urgent_intake is not None:
        intake_orders: list[UrgentOrderDraft] = []
        for item in urgent_intake.get("orders", []):
            draft_data = dict(item)
            if draft_data.get("priority") is None:
                draft_data.pop("priority", None)
            intake_orders.append(UrgentOrderDraft.model_validate(draft_data))
        intake_understanding = UrgentUnderstanding(
            is_urgent_insertion=True,
            action=urgent_intake["action"],
            orders=intake_orders,
            referenced_order_ids=[
                str(item) for item in urgent_intake.get("referenced_order_ids", [])
            ],
        )
        return finish_urgent_workflow(
            intake_understanding,
            result,
            entry_tool="begin_urgent_insertion",
        )

    # A plan requested through the Agent is persisted here, after the SDK has
    # selected and executed plan_dispatch.  This keeps the conversation as the
    # orchestration entry point while preserving the same immutable plan store
    # used by the REST endpoints.
    plan_tool_used = any(item.get("tool") == "plan_dispatch" for item in context.evidence)
    if (
        record is None
        and plan_tool_used
        and dataset_record is not None
        and context.plan is not None
    ):
        validation = validate_plan(dataset_record.dataset, context.plan, matrix)
        plan_id = f"PLAN-{uuid4().hex[:12].upper()}"
        record = PlanRecord(
            plan_id=plan_id,
            dataset_id=dataset_record.dataset_id,
            version=1,
            state="PROPOSED",
            plan=context.plan,
            validation=validation,
            matrix=matrix,
            created_at=datetime.now(UTC).isoformat(),
        )
        store.add_plan(record)
        repository.save_plan(record)
        session.plan_id = plan_id
        session.plan_version = 1
    if (
        record is not None
        and agent_dataset_record is not None
        and context.pending_preview_plan is not None
    ):
        option = _persist_plan_change_preview(
            context,
            preview_record_for_change or record,
            agent_dataset_record,
        )
        if option is not None and context.evidence:
            options = [option]
            if context.pending_preview_kind == "PRIORITIZE_ORDER":
                base_for_keep = preview_record_for_change or record
                base_summary = _plan_payload(base_for_keep)["summary"]
                options.append(
                    {
                        **option,
                        "option_id": f"{option['option_id']}-KEEP",
                        "label": "方案 B",
                        "title": "維持原順序",
                        "rationale": "維持目前剩餘站點順序；需回覆客戶送不到。",
                        "preview_version": base_for_keep.version,
                        "feasible": True,
                        "selectable": False,
                        "requires_human_confirmation": False,
                        "cost": {
                            "distance_delta_m": 0,
                            "distance_delta_km": 0,
                            "duration_delta_s": 0,
                            "duration_delta_min": 0,
                            "vehicle_change_count": 0,
                            "minimum_capacity_slack_kg": None,
                        },
                        "affected_vehicle_count": 0,
                        "moved_order_count": 0,
                        "reordered_order_count": 0,
                        "after": base_summary,
                        "diff": compute_plan_diff(base_for_keep.plan, base_for_keep.plan),
                    }
                )
            context.evidence[-1].update(
                {
                    "options": options,
                    "preview_version": option["preview_version"],
                    "before": option["before"],
                    "after": option["after"],
                    "diff": option["diff"],
                }
            )
    evidence = []
    for item in context.evidence:
        evidence.append(
            {
                "tool": item.get("tool", "unknown"),
                "data": {key: value for key, value in item.items() if key != "tool"},
            }
        )
    requires_confirmation = bool(record and record.state in {"DRAFT", "VALIDATED", "PROPOSED"})
    usage = {
        "total_tokens": context.budget.total_tokens,
        "tool_calls": context.budget.tool_calls,
        "agent_run_id": context.agent_run_id,
    }
    session.history.extend(
        [
            ("user", _safe_session_text(payload.message)),
            ("assistant", _safe_session_text(final_output)),
        ]
    )
    session.last_tool = context.evidence[-1].get("tool") if context.evidence else None
    session.rule_source_utterance = context.rule_source_utterance
    session.pending_fields = tuple(context.pending_fields)
    session.frozen_stop_ids = tuple(context.frozen_stop_ids)
    session.frozen_stop_count = len(session.frozen_stop_ids)
    for evidence_item in context.evidence:
        evidence_order_id = evidence_item.get("order_id")
        if isinstance(evidence_order_id, str):
            session.order_id = evidence_order_id
        evidence_vehicle_id = evidence_item.get("vehicle_id")
        if not isinstance(evidence_vehicle_id, str):
            evidence_vehicle_id = evidence_item.get("target_vehicle_id")
        if isinstance(evidence_vehicle_id, str):
            session.vehicle_id = evidence_vehicle_id
        if evidence_item.get("objective") in {"FASTEST", "BALANCED", "STABLE"}:
            session.strategy = str(evidence_item["objective"])
        preview_version = evidence_item.get("preview_version")
        if isinstance(preview_version, int):
            session.last_preview_version = preview_version
        validation = evidence_item.get("validation")
        if isinstance(validation, dict):
            errors = validation.get("errors")
            if isinstance(errors, list):
                session.pending_fields = tuple(
                    str(item.get("path"))
                    for item in errors
                    if isinstance(item, dict) and isinstance(item.get("path"), str)
                )
    if len(session.history) > 12:
        session.history = session.history[-12:]
    repository.save_agent_session(
        payload.session_id,
        _session_payload(session),
        datetime.now(UTC).isoformat(),
    )
    return {
        "session_id": payload.session_id,
        "agent_run_id": context.agent_run_id,
        "message": final_output,
        "evidence": evidence,
        "requires_human_confirmation": requires_confirmation,
        "usage": usage,
        "provider_mode": record.matrix.provider_mode if record else matrix.provider_mode,
        "plan_id": record.plan_id if record else None,
        "plan_version": record.version if record else None,
        "runner_result_type": type(result).__name__,
        "request_id": _request_id(request),
    }


# In the Render image the Vite build is copied to /app/frontend/dist. Keeping
# the same path convention locally makes the production container and local
# smoke tests exercise the same SPA fallback behaviour.
_frontend_dist = Path(__file__).resolve().parents[2] / "frontend" / "dist"
if _frontend_dist.is_dir():
    app.mount(
        "/assets",
        StaticFiles(directory=_frontend_dist / "assets"),
        name="frontend-assets",
    )

    @app.get("/{full_path:path}", include_in_schema=False)
    async def spa_fallback(full_path: str) -> Any:
        requested = (_frontend_dist / full_path).resolve()
        try:
            requested.relative_to(_frontend_dist.resolve())
        except ValueError:
            return JSONResponse(status_code=404, content={"detail": "Not found"})
        if requested.is_file():
            return FileResponse(requested)
        index_file = _frontend_dist / "index.html"
        if index_file.is_file():
            return FileResponse(index_file)
        return JSONResponse(status_code=404, content={"detail": "Frontend build unavailable"})
